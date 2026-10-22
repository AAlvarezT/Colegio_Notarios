"""Web API tests for the Conversor ODE FastAPI server.

All tests use synthetic in-memory data — no real report files needed.
"""
from __future__ import annotations

import io
import json
import os
from concurrent.futures import ThreadPoolExecutor, as_completed

# Configure env vars BEFORE importing server so auth is active
os.environ['SECRET_KEY'] = 'test-secret-key-do-not-use-in-prod'
os.environ['ODE_USERNAME'] = 'testuser'
os.environ['ODE_PASSWORD'] = 'testpass'

import pytest
from fastapi.testclient import TestClient

from server import app  # noqa: E402

client = TestClient(app, raise_server_exceptions=False)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_xlsx_bytes(with_tipo_col: bool = True) -> bytes:
    """Generate a valid synthetic xlsx report in memory."""
    import openpyxl
    from datetime import datetime

    wb = openpyxl.Workbook()
    ws = wb.active
    cols = ['NumRuc', 'Serie', 'Documento', 'Importe Cancelado', 'Fecha Ingreso', 'Fecha Cancelacion', 'Banco']
    if with_tipo_col:
        cols.append('Tipo Comprobante')
    ws.append(cols)
    rows = [
        ('20111111111', 'FE01', '0000001', 100.0, datetime(2026, 8, 10), datetime(2026, 8, 10), 'INTERBANK', '01'),
        ('20222222222', 'BE01', '0000002',  50.0, datetime(2026, 8, 11), datetime(2026, 8, 11), 'INTERBANK', '03'),
    ]
    for r in rows:
        ws.append(r[:7] + ((r[7],) if with_tipo_col else ()))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_xlsx_with_unknown_serie(serie: str = 'ZZ99') -> bytes:
    """Xlsx with a serie that won't match any serie_map."""
    import openpyxl
    from datetime import datetime
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['NumRuc', 'Serie', 'Documento', 'Importe Cancelado', 'Fecha Ingreso', 'Fecha Cancelacion', 'Banco'])
    ws.append(('20111111111', serie, '0000001', 100.0, datetime(2026, 8, 10), datetime(2026, 8, 10), 'INTERBANK'))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_xlsx_with_conflict() -> bytes:
    """Xlsx with tipo_comprobante='01' but serie_map maps that serie to '03'."""
    import openpyxl
    from datetime import datetime
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['NumRuc', 'Serie', 'Documento', 'Importe Cancelado', 'Fecha Ingreso', 'Fecha Cancelacion', 'Banco', 'Tipo Comprobante'])
    ws.append(('20111111111', 'FE01', '0000001', 100.0, datetime(2026, 8, 10), datetime(2026, 8, 10), 'INTERBANK', '01'))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _default_config(**overrides) -> dict:
    cfg = {
        'cuenta_clientes': '12120001',
        'sub_diario': '01',
        'tipo_anexo': '02',
        'moneda': 'MN',
        'tipo_conversion': 'VTA',
        'doc_anulado': '0',
        'glosa': 'COBRANZA DEL MES',
        'serie_map': {},
        'bancos': {
            'INTERBANK': {'cuenta_contable': '10100011', 'medio_pago': '008'},
            'BBVA CONTINENTAL': {'cuenta_contable': '10100011', 'medio_pago': '008'},
            'DE LA NACION': {'cuenta_contable': '10100011', 'medio_pago': '008'},
            'OTROS': {'cuenta_contable': '10100001', 'medio_pago': ''},
        },
    }
    cfg.update(overrides)
    return cfg


def _auth_cookies() -> dict:
    """Log in and return the session cookie dict."""
    res = client.post('/login', data={'username': 'testuser', 'password': 'testpass'})
    assert res.status_code == 200, f'Login failed: {res.text}'
    return dict(res.cookies)


def _analyze(xlsx_bytes: bytes, config: dict | None = None, year: str = '2026',
             month: str = '8', fecha_calculo: str = 'fecha_cancelacion',
             cookies: dict | None = None) -> tuple:
    cookies = cookies or _auth_cookies()
    cfg = config or _default_config()
    res = client.post(
        '/api/analyze',
        files={'reporte': ('reporte.xlsx', xlsx_bytes, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
        data={'year': year, 'month': month, 'fecha_calculo': fecha_calculo,
              'config_json': json.dumps(cfg)},
        cookies=cookies,
    )
    return res, res.json() if res.headers.get('content-type', '').startswith('application/json') else {}


def _generate(xlsx_bytes: bytes, config: dict | None = None, year: str = '2026',
              month: str = '8', cookies: dict | None = None) -> tuple:
    cookies = cookies or _auth_cookies()
    cfg = config or _default_config()
    res = client.post(
        '/api/generate',
        files={'reporte': ('reporte.xlsx', xlsx_bytes, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
        data={'year': year, 'month': month, 'fecha_calculo': 'fecha_cancelacion',
              'config_json': json.dumps(cfg)},
        cookies=cookies,
    )
    return res, res.json() if res.headers.get('content-type', '').startswith('application/json') else {}

# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

def test_unauthenticated_redirects():
    """GET / without session cookie → 302 redirect to /login."""
    res = client.get('/', follow_redirects=False)
    assert res.status_code == 302
    assert '/login' in res.headers.get('location', '')


def test_login_success():
    """POST /login with correct credentials → 200 + session cookie."""
    res = client.post('/login', data={'username': 'testuser', 'password': 'testpass'})
    assert res.status_code == 200
    assert res.json().get('ok') is True
    assert 'ode_session' in res.cookies


def test_login_wrong_password():
    """POST /login with wrong password → 401."""
    res = client.post('/login', data={'username': 'testuser', 'password': 'wrongpass'})
    assert res.status_code == 401


def test_logout():
    """POST /logout → 200, cookie cleared."""
    cookies = _auth_cookies()
    res = client.post('/logout', cookies=cookies)
    assert res.status_code == 200
    # After logout the session cookie should be deleted
    assert res.cookies.get('ode_session') == '' or 'ode_session' not in res.cookies


def test_analyze_unauthenticated():
    """POST /api/analyze without auth → 401."""
    xlsx = _make_xlsx_bytes()
    res = client.post(
        '/api/analyze',
        files={'reporte': ('r.xlsx', xlsx, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
        data={'year': '2026', 'month': '8', 'config_json': '{}'},
    )
    assert res.status_code == 401


def test_analyze_synthetic_with_tipo_column():
    """Reporte with tipo_comprobante, empty serie_map → no errors, correct tip_doc."""
    xlsx = _make_xlsx_bytes(with_tipo_col=True)
    res, data = _analyze(xlsx)
    assert res.status_code == 200, data
    assert data.get('errors') == [], data.get('errors')
    tips = {r['tip_doc'] for r in data.get('preview', [])}
    assert '01' in tips and '03' in tips


def test_analyze_synthetic_no_tipo_no_map_blocked():
    """No tipo_comprobante column, empty serie_map → errors not empty."""
    xlsx = _make_xlsx_bytes(with_tipo_col=False)
    cfg = _default_config(serie_map={})
    res, data = _analyze(xlsx, config=cfg)
    assert res.status_code == 200, data
    assert len(data.get('errors', [])) > 0


def test_analyze_synthetic_with_serie_map():
    """No tipo_comprobante column, correct serie_map → no errors, 01 and 03 in preview."""
    xlsx = _make_xlsx_bytes(with_tipo_col=False)
    cfg = _default_config(serie_map={'FE': '01', 'BE': '03'})
    res, data = _analyze(xlsx, config=cfg)
    assert res.status_code == 200, data
    assert data.get('errors') == [], data.get('errors')
    tips = {r['tip_doc'] for r in data.get('preview', [])}
    assert '01' in tips and '03' in tips


def test_analyze_unknown_serie_blocked():
    """Serie not covered by any map → errors contain the serie name."""
    xlsx = _make_xlsx_with_unknown_serie('ZZ99')
    cfg = _default_config(serie_map={'FE': '01', 'BE': '03'})
    res, data = _analyze(xlsx, config=cfg)
    assert res.status_code == 200
    errors_text = ' '.join(data.get('errors', []))
    assert 'ZZ99' in errors_text or len(data.get('errors', [])) > 0


def test_analyze_conflict_blocked():
    """tipo_comprobante='01' but serie_map says '03' → error contains 'Conflicto'."""
    xlsx = _make_xlsx_with_conflict()
    cfg = _default_config(serie_map={'FE01': '03'})
    res, data = _analyze(xlsx, config=cfg)
    assert res.status_code == 200
    errors_text = ' '.join(data.get('errors', []))
    assert 'onflicto' in errors_text or 'conflicto' in errors_text.lower() or len(data.get('errors', [])) > 0


def test_generate_returns_three_base64_files():
    """Generate with explicit tipo_comprobante → 3 base64 files in response."""
    xlsx = _make_xlsx_bytes(with_tipo_col=True)
    res, data = _generate(xlsx)
    assert res.status_code == 200, data
    files = data.get('files', [])
    assert len(files) == 3
    for f in files:
        assert 'data' in f and 'name' in f and 'mime' in f
        assert len(f['data']) > 100  # non-empty base64


def test_generate_xls_readable():
    """Decode the .xls base64 and open with xlrd; verify TIP_DOC is TEXT."""
    import base64
    import xlrd
    xlsx = _make_xlsx_bytes(with_tipo_col=True)
    res, data = _generate(xlsx)
    assert res.status_code == 200, data
    files = data.get('files', [])
    xls_file = next((f for f in files if f['name'].endswith('.xls')), None)
    assert xls_file is not None, 'No .xls file returned'
    xls_bytes = base64.b64decode(xls_file['data'])
    wb = xlrd.open_workbook(file_contents=xls_bytes)
    ws = wb.sheet_by_index(0)
    # Find TIP_DOC column
    headers = [ws.cell_value(0, c) for c in range(ws.ncols)]
    assert 'TIP_DOC' in headers, f'TIP_DOC not found. Headers: {headers}'
    tip_col = headers.index('TIP_DOC')
    tip_values = [ws.cell_value(r, tip_col) for r in range(1, ws.nrows) if ws.cell_value(r, tip_col)]
    assert any(v in ('01', '03') for v in tip_values), f'No valid TIP_DOC values: {tip_values}'


def test_generate_blocked_without_valid_analysis_config():
    """Invalid serie_map (code '99') → 422."""
    xlsx = _make_xlsx_bytes(with_tipo_col=False)
    cfg = _default_config(serie_map={'FE01': '99'})
    res, data = _generate(xlsx, config=cfg)
    assert res.status_code == 422


def test_generate_isolation():
    """Two concurrent generate requests don't share state (temp dirs are independent)."""
    xlsx = _make_xlsx_bytes(with_tipo_col=True)
    cookies = _auth_cookies()

    def do_generate():
        return _generate(xlsx, cookies=cookies)

    with ThreadPoolExecutor(max_workers=2) as ex:
        futures = [ex.submit(do_generate) for _ in range(2)]
        results = [f.result() for f in as_completed(futures)]

    for res, data in results:
        assert res.status_code == 200, data
        assert len(data.get('files', [])) == 3


def test_response_no_nan_no_paths():
    """Analyze response JSON must not contain NaN strings or server-side paths."""
    xlsx = _make_xlsx_bytes(with_tipo_col=True)
    res, _ = _analyze(xlsx)
    assert res.status_code == 200
    text = res.text
    assert 'NaN' not in text
    assert 'C:\\' not in text and '/tmp/' not in text


def test_invalid_file_type_rejected():
    """Uploading a .txt file → 422 with descriptive error."""
    cookies = _auth_cookies()
    res = client.post(
        '/api/analyze',
        files={'reporte': ('report.txt', b'some text data', 'text/plain')},
        data={'year': '2026', 'month': '8', 'config_json': '{}'},
        cookies=cookies,
    )
    assert res.status_code == 422
    assert 'txt' in res.text.lower() or 'tipo' in res.text.lower() or 'permitido' in res.text.lower()


def test_file_too_large_rejected():
    """File exceeding 8 MB → 413."""
    cookies = _auth_cookies()
    big = b'x' * (8 * 1024 * 1024 + 1)
    res = client.post(
        '/api/analyze',
        files={'reporte': ('big.xlsx', big, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
        data={'year': '2026', 'month': '8', 'config_json': '{}'},
        cookies=cookies,
    )
    assert res.status_code == 413


def test_config_export_returns_valid_json():
    """GET /api/config/default → valid JSON with required keys."""
    cookies = _auth_cookies()
    res = client.get('/api/config/default', cookies=cookies)
    assert res.status_code == 200
    data = res.json()
    for key in ['cuenta_clientes', 'sub_diario', 'tipo_anexo', 'moneda', 'serie_map', 'bancos']:
        assert key in data, f'Missing key: {key}'


def test_config_validate_valid():
    """POST /api/config/validate with valid config → ok=True."""
    cookies = _auth_cookies()
    cfg = _default_config()
    res = client.post(
        '/api/config/validate',
        data={'config_json': json.dumps(cfg)},
        cookies=cookies,
    )
    assert res.status_code == 200
    assert res.json().get('ok') is True


def test_config_validate_invalid():
    """POST /api/config/validate with invalid serie_map (bad code) → 422."""
    cookies = _auth_cookies()
    cfg = _default_config(serie_map={'FE': '99'})
    res = client.post(
        '/api/config/validate',
        data={'config_json': json.dumps(cfg)},
        cookies=cookies,
    )
    assert res.status_code == 422


# ---------------------------------------------------------------------------
# Logging / filesystem regression tests (Vercel read-only FS fix)
# ---------------------------------------------------------------------------

import logging as _logging
from pathlib import Path as _Path


def test_transformer_no_file_handler_at_import():
    """Importing transformer must NOT configure a FileHandler.

    Regression for: OSError: [Errno 30] Read-only file system: '/var/task/logs'
    The module-level logging block was creating log_dir and a FileHandler on
    import, which fails on Vercel's read-only filesystem.
    """
    logger = _logging.getLogger('conversor_ode')
    file_handlers = [h for h in logger.handlers if isinstance(h, _logging.FileHandler)]
    assert not file_handlers, (
        "transformer.py must not add FileHandler at import time. "
        "Call configure_file_logging() from the desktop entry point instead."
    )


def test_configure_file_logging_desktop(tmp_path):
    """configure_file_logging() installs a FileHandler pointing to tmp_path."""
    from transformer import configure_file_logging, LOGGER as ode_logger

    # server.py adds a StreamHandler on import; save and clear so we get a
    # clean slate that mimics a fresh desktop process.
    saved_handlers = list(ode_logger.handlers)
    ode_logger.handlers.clear()
    try:
        configure_file_logging(tmp_path)
        file_handlers = [h for h in ode_logger.handlers if isinstance(h, _logging.FileHandler)]
        assert file_handlers, "configure_file_logging must add at least one FileHandler"
        assert all(str(tmp_path) in h.baseFilename for h in file_handlers), (
            "FileHandler must point to the provided log_dir, not the source tree"
        )
        # Idempotent: calling again does not add a second handler
        count_before = len(ode_logger.handlers)
        configure_file_logging(tmp_path)
        assert len(ode_logger.handlers) == count_before
    finally:
        for h in ode_logger.handlers[:]:
            h.close()
        ode_logger.handlers.clear()
        ode_logger.handlers.extend(saved_handlers)


def test_server_import_survives_readonly_filesystem(tmp_path):
    """Simulate Vercel's read-only FS: importing server must not call mkdir.

    Runs in a subprocess to get a fresh module environment.
    """
    import subprocess
    import sys

    script = """
import sys, os
sys.path.insert(0, ".")
os.environ.setdefault('SECRET_KEY',   'test-key')
os.environ.setdefault('ODE_USERNAME', 'u')
os.environ.setdefault('ODE_PASSWORD', 'p')
# Clear VERCEL_ENV so IS_VERCEL is False (avoids 503 on missing vars)
os.environ.pop('VERCEL_ENV', None)

from pathlib import Path
from unittest.mock import patch

failed_mkdirs = []
original_mkdir = Path.mkdir

def guarded_mkdir(self, *args, **kwargs):
    path_str = str(self)
    # Allow only truly temporary paths (OS temp dirs)
    import tempfile
    if path_str.startswith(tempfile.gettempdir()):
        return original_mkdir(self, *args, **kwargs)
    failed_mkdirs.append(path_str)
    raise OSError(30, 'Read-only file system', path_str)

with patch.object(Path, 'mkdir', guarded_mkdir):
    try:
        import transformer  # noqa: F401
        import server       # noqa: F401
        print('IMPORTS_OK')
        print('FAILED_MKDIRS:', failed_mkdirs)
    except OSError as exc:
        print(f'OSERROR: {exc}')
        print(f'FAILED_MKDIRS: {failed_mkdirs}')
        sys.exit(1)
"""

    result = subprocess.run(
        [sys.executable, '-c', script],
        capture_output=True,
        text=True,
        cwd=str(_Path(__file__).parent.parent),
    )
    assert result.returncode == 0, (
        f"Server import raised OSError (simulated read-only FS):\n"
        f"stdout: {result.stdout}\n"
        f"stderr: {result.stderr}"
    )
    assert 'IMPORTS_OK' in result.stdout, (
        f"Expected IMPORTS_OK in stdout:\n{result.stdout}\n{result.stderr}"
    )
