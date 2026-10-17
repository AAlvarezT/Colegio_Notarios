"""
Tests for mixed-type (Factura 01 + Boleta 03) support.

All scenarios use synthetic DataFrames; no real file is required
unless the test explicitly loads SOURCE_PATH.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
import xlrd

from transformer import (
    DEFAULT_CONFIG,
    OUTPUT_COLUMNS,
    build_carga_dataframe,
    build_conciliacion,
    export_carga_xls,
    normalize_tip_doc_code,
    process_report,
    resolve_tip_doc,
    resolve_tip_doc_for_df,
    split_ode_rows,
    validate_report,
    validate_serie_map,
)

_DATE = datetime(2026, 8, 15)


def _make_df(**overrides) -> pd.DataFrame:
    """Minimal valid processed DataFrame with two rows."""
    base = {
        'numruc': ['20000000001', '20000000002'],
        'serie': ['FE01', 'BE01'],
        'documento': ['0000001', '0000002'],
        'importe_cancelado': [100.0, 50.0],
        'banco': ['INTERBANK', 'INTERBANK'],
        'fecha_cancelacion': [_DATE, _DATE],
    }
    base.update(overrides)
    return pd.DataFrame(base)


def _cfg_serie_map(**kw) -> dict:
    return {**DEFAULT_CONFIG, 'serie_map': {'FE': '01', 'BE': '03', **kw}}


# ---------------------------------------------------------------------------
# 1. normalize_tip_doc_code — strict validation
# ---------------------------------------------------------------------------

class TestNormalizeTipDocCode:
    # Accepted values
    def test_string_01(self):        assert normalize_tip_doc_code('01')  == '01'
    def test_string_03(self):        assert normalize_tip_doc_code('03')  == '03'
    def test_int_1(self):            assert normalize_tip_doc_code(1)     == '01'
    def test_int_3(self):            assert normalize_tip_doc_code(3)     == '03'
    def test_float_1_0(self):        assert normalize_tip_doc_code(1.0)   == '01'
    def test_float_3_0(self):        assert normalize_tip_doc_code(3.0)   == '03'
    def test_bare_string_1(self):    assert normalize_tip_doc_code('1')   == '01'
    def test_bare_string_3(self):    assert normalize_tip_doc_code('3')   == '03'
    def test_np_int64_1(self):       assert normalize_tip_doc_code(np.int64(1)) == '01'
    def test_np_int64_3(self):       assert normalize_tip_doc_code(np.int64(3)) == '03'

    # Rejected — unknown code
    def test_unknown_code_02(self):  assert normalize_tip_doc_code('02')  is None
    def test_unknown_str(self):      assert normalize_tip_doc_code('FAC') is None
    def test_empty_str(self):        assert normalize_tip_doc_code('')    is None
    def test_none(self):             assert normalize_tip_doc_code(None)  is None

    # Rejected — NaN / Inf
    def test_float_nan(self):        assert normalize_tip_doc_code(float('nan'))  is None
    def test_np_nan(self):           assert normalize_tip_doc_code(np.nan)        is None
    def test_float_inf(self):        assert normalize_tip_doc_code(float('inf'))  is None
    def test_float_neg_inf(self):    assert normalize_tip_doc_code(float('-inf')) is None
    def test_np_inf(self):           assert normalize_tip_doc_code(np.inf)        is None

    # Rejected — non-integer floats (no truncation/rounding allowed)
    def test_float_1_9_rejected(self):
        """1.9 must NOT truncate to 1 and return '01'."""
        assert normalize_tip_doc_code(1.9) is None

    def test_float_3_7_rejected(self):
        assert normalize_tip_doc_code(3.7) is None

    def test_float_1_1_rejected(self):
        assert normalize_tip_doc_code(1.1) is None

    # Rejected — booleans (bool subclasses int in Python)
    def test_bool_true_rejected(self):
        assert normalize_tip_doc_code(True) is None

    def test_bool_false_rejected(self):
        assert normalize_tip_doc_code(False) is None

    def test_np_bool_rejected(self):
        assert normalize_tip_doc_code(np.bool_(True)) is None


# ---------------------------------------------------------------------------
# 2. validate_serie_map
# ---------------------------------------------------------------------------

class TestValidateSerieMap:
    def test_empty_is_valid(self):
        assert validate_serie_map({}) is None

    def test_valid_map(self):
        assert validate_serie_map({'FE': '01', 'BE': '03'}) is None

    def test_empty_key_blocked(self):
        msg = validate_serie_map({'': '01', 'BE': '03'})
        assert msg is not None and 'vac' in msg.lower()

    def test_whitespace_key_blocked(self):
        msg = validate_serie_map({'   ': '01'})
        assert msg is not None

    def test_contradictory_normalised_keys_blocked(self):
        """'FE' and 'fe' normalise to the same key but have different types."""
        msg = validate_serie_map({'FE': '01', 'fe': '03'})
        assert msg is not None and 'contradictorias' in msg.lower()

    def test_same_normalised_key_same_type_ok(self):
        assert validate_serie_map({'FE': '01', 'fe': '01'}) is None

    def test_spaces_in_keys_normalised_and_contradictory(self):
        msg = validate_serie_map({' FE ': '01', 'FE ': '03'})
        assert msg is not None


# ---------------------------------------------------------------------------
# 3. resolve_tip_doc — per-row logic
# ---------------------------------------------------------------------------

class TestResolveTipDoc:
    # 3a. Empty map + no explicit column → UNRESOLVED (NO global fallback)
    def test_empty_map_no_explicit_unresolved(self):
        cfg = {**DEFAULT_CONFIG, 'tipo_documento': '01', 'serie_map': {}}
        td, msg = resolve_tip_doc({'serie': 'FE01'}, cfg)
        assert td is None and msg is not None

    def test_empty_map_boleta_series_also_unresolved(self):
        cfg = {**DEFAULT_CONFIG, 'tipo_documento': '03', 'serie_map': {}}
        td, msg = resolve_tip_doc({'serie': 'BE01'}, cfg)
        assert td is None and msg is not None

    # 3b. Exact key match
    def test_exact_match_factura(self):
        cfg = {**DEFAULT_CONFIG, 'serie_map': {'FE01': '01'}}
        td, msg = resolve_tip_doc({'serie': 'FE01'}, cfg)
        assert td == '01' and msg is None

    def test_exact_match_boleta(self):
        cfg = {**DEFAULT_CONFIG, 'serie_map': {'BE01': '03'}}
        td, msg = resolve_tip_doc({'serie': 'BE01'}, cfg)
        assert td == '03' and msg is None

    # 3c. Prefix match
    def test_prefix_match_factura(self):
        td, msg = resolve_tip_doc({'serie': 'FE01'}, _cfg_serie_map())
        assert td == '01' and msg is None

    def test_prefix_match_boleta(self):
        td, msg = resolve_tip_doc({'serie': 'BE01'}, _cfg_serie_map())
        assert td == '03' and msg is None

    # 3d. Exact match beats shorter prefix
    def test_exact_over_shorter_prefix(self):
        """{'F': '01', 'FE01': '03'} → exact key 'FE01' wins → '03'."""
        cfg = {**DEFAULT_CONFIG, 'serie_map': {'F': '01', 'FE01': '03'}}
        td, _ = resolve_tip_doc({'serie': 'FE01'}, cfg)
        assert td == '03'

    # 3e. Longest prefix beats shorter prefix
    def test_longest_prefix_wins(self):
        """{'F': '01', 'FE': '03'} → 'FE' (len 2) beats 'F' (len 1) → '03'."""
        cfg = {**DEFAULT_CONFIG, 'serie_map': {'F': '01', 'FE': '03'}}
        td, _ = resolve_tip_doc({'serie': 'FE01'}, cfg)
        assert td == '03'

    # 3f. Order-independent result
    def test_longest_prefix_order_independent(self):
        cfg_1 = {**DEFAULT_CONFIG, 'serie_map': {'FE': '03', 'F': '01'}}
        cfg_2 = {**DEFAULT_CONFIG, 'serie_map': {'F': '01', 'FE': '03'}}
        td1, _ = resolve_tip_doc({'serie': 'FE01'}, cfg_1)
        td2, _ = resolve_tip_doc({'serie': 'FE01'}, cfg_2)
        assert td1 == td2 == '03'

    # 3g. Explicit tipo_comprobante column
    def test_explicit_column_01(self):
        td, msg = resolve_tip_doc({'serie': 'X99', 'tipo_comprobante': '01'}, DEFAULT_CONFIG)
        assert td == '01' and msg is None

    def test_explicit_column_03(self):
        td, msg = resolve_tip_doc({'serie': 'X99', 'tipo_comprobante': '03'}, DEFAULT_CONFIG)
        assert td == '03' and msg is None

    def test_explicit_column_integer_3(self):
        td, msg = resolve_tip_doc({'serie': 'X99', 'tipo_comprobante': 3}, DEFAULT_CONFIG)
        assert td == '03' and msg is None

    # 3h. Unknown/invalid → block
    def test_unknown_explicit_blocks(self):
        td, msg = resolve_tip_doc({'serie': 'X99', 'tipo_comprobante': '99'}, DEFAULT_CONFIG)
        assert td is None and '99' in msg

    def test_unknown_serie_with_map_blocks(self):
        cfg = {**DEFAULT_CONFIG, 'serie_map': {'FE': '01'}}
        td, msg = resolve_tip_doc({'serie': 'BE01'}, cfg)
        assert td is None and ('BE01' in msg or 'BE' in msg)

    # 3i. Conflict between explicit column and map → block
    def test_conflict_explicit_vs_map_blocks(self):
        cfg = {**DEFAULT_CONFIG, 'serie_map': {'FE': '01'}}
        td, msg = resolve_tip_doc({'serie': 'FE01', 'tipo_comprobante': '03'}, cfg)
        assert td is None and 'Conflicto' in msg

    # 3j. Invalid code in map → block
    def test_invalid_code_in_map_blocks(self):
        cfg = {**DEFAULT_CONFIG, 'serie_map': {'FE': '99'}}
        td, msg = resolve_tip_doc({'serie': 'FE01'}, cfg)
        assert td is None and '99' in msg


# ---------------------------------------------------------------------------
# 4. resolve_tip_doc_for_df
# ---------------------------------------------------------------------------

class TestResolveTipDocForDf:
    def test_all_resolved_via_map(self):
        df = _make_df()
        series, obs = resolve_tip_doc_for_df(df, _cfg_serie_map())
        assert list(series) == ['01', '03']
        assert obs == []

    def test_empty_map_all_rows_unresolved(self):
        """Empty serie_map + no explicit column → every row unresolved."""
        df = _make_df()
        series, obs = resolve_tip_doc_for_df(df, {**DEFAULT_CONFIG, 'serie_map': {}})
        assert list(series) == ['', '']
        assert len(obs) == 2

    def test_partial_unresolved(self):
        df = _make_df(serie=['FE01', 'ZZ99'])
        cfg = {**DEFAULT_CONFIG, 'serie_map': {'FE': '01'}}
        series, obs = resolve_tip_doc_for_df(df, cfg)
        assert series.iloc[0] == '01'
        assert series.iloc[1] == ''
        assert len(obs) == 1 and obs[0]['serie'] == 'ZZ99'

    def test_conflict_marked_correctly(self):
        df = _make_df(serie=['FE01', 'FE01'], tipo_comprobante=['01', '03'])
        cfg = {**DEFAULT_CONFIG, 'serie_map': {'FE': '01'}}
        series, obs = resolve_tip_doc_for_df(df, cfg)
        assert series.iloc[0] == '01'
        assert series.iloc[1] == ''
        assert any(o['tipo'] == 'CONFLICTO' or 'Conflicto' in o['mensaje'] for o in obs)

    def test_invalid_map_marks_all_rows_error_mapa(self):
        df = _make_df()
        cfg = {**DEFAULT_CONFIG, 'serie_map': {'FE': '01', 'fe': '03'}}
        series, obs = resolve_tip_doc_for_df(df, cfg)
        assert all(s == '' for s in series)
        assert all(o['tipo'] == 'ERROR_MAPA' for o in obs)


# ---------------------------------------------------------------------------
# 5. validate_report — TIP_DOC blocking
# ---------------------------------------------------------------------------

class TestValidateReportTipDoc:
    def test_passes_with_valid_map(self):
        df = _make_df()
        errors, _, _ = validate_report(df, periodo='202608', config=_cfg_serie_map())
        assert errors == []

    def test_empty_map_blocks_every_row(self):
        df = _make_df()
        errors, _, _ = validate_report(df, periodo='202608', config={**DEFAULT_CONFIG, 'serie_map': {}})
        assert len(errors) >= 2

    def test_blocks_on_unresolved_serie(self):
        df = _make_df(serie=['FE01', 'ZZ99'])
        cfg = {**DEFAULT_CONFIG, 'serie_map': {'FE': '01'}}
        errors, _, _ = validate_report(df, periodo='202608', config=cfg)
        assert any('ZZ99' in e for e in errors)

    def test_blocks_on_conflict(self):
        df = _make_df(serie=['FE01', 'FE01'], tipo_comprobante=['01', '03'])
        cfg = {**DEFAULT_CONFIG, 'serie_map': {'FE': '01'}}
        errors, _, _ = validate_report(df, periodo='202608', config=cfg)
        assert any('Conflicto' in e or 'conflicto' in e for e in errors)

    def test_no_config_skips_tip_doc_check(self):
        """config=None → TIP_DOC check must be skipped entirely."""
        df = _make_df(serie=['ZZ99', 'ZZ88'])
        errors, _, _ = validate_report(df, periodo='202608')
        assert not any('tipo_comprobante' in e.lower() or 'tip_doc' in e.lower() for e in errors)


# ---------------------------------------------------------------------------
# 6. build_carga_dataframe — mixed types
# ---------------------------------------------------------------------------

class TestBuildCargaMixed:
    def _build(self, **overrides):
        return build_carga_dataframe(_make_df(**overrides), periodo='202608', config=_cfg_serie_map())

    def test_h_rows_correct_tip_doc(self):
        carga = self._build()
        h = carga[carga['DEBE HABER'] == 'H']
        assert list(h['TIP_DOC']) == ['01', '03']

    def test_d_rows_tip_doc_blank(self):
        carga = self._build()
        d = carga[carga['DEBE HABER'] == 'D']
        assert d['TIP_DOC'].astype(str).str.strip().eq('').all()

    def test_accounting_balance(self):
        carga = self._build()
        d = carga[carga['DEBE HABER'] == 'D']['IMPORT_TOTAL'].sum()
        h = carga[carga['DEBE HABER'] == 'H']['IMPORT_TOTAL'].sum()
        assert abs(d - h) <= 0.01

    def test_raises_on_unresolved_serie(self):
        df = _make_df(serie=['FE01', 'ZZ99'])
        cfg = {**DEFAULT_CONFIG, 'serie_map': {'FE': '01'}}
        with pytest.raises(ValueError, match='ZZ99'):
            build_carga_dataframe(df, periodo='202608', config=cfg)

    def test_empty_map_no_explicit_raises(self):
        """Empty map + no explicit column → MUST fail (no silent fallback)."""
        df = _make_df()
        cfg = {**DEFAULT_CONFIG, 'tipo_documento': '01', 'serie_map': {}}
        with pytest.raises(ValueError):
            build_carga_dataframe(df, periodo='202608', config=cfg)

    def test_explicit_column_with_empty_map(self):
        """Explicit tipo_comprobante column works even with empty serie_map."""
        df = _make_df(tipo_comprobante=['01', '03'])
        cfg = {**DEFAULT_CONFIG, 'serie_map': {}}
        carga = build_carga_dataframe(df, periodo='202608', config=cfg)
        h = carga[carga['DEBE HABER'] == 'H']
        assert list(h['TIP_DOC']) == ['01', '03']

    def test_invalid_serie_map_raises_before_processing(self):
        cfg = {**DEFAULT_CONFIG, 'serie_map': {'': '01'}}
        with pytest.raises(ValueError, match='serie_map'):
            build_carga_dataframe(_make_df(), periodo='202608', config=cfg)


# ---------------------------------------------------------------------------
# 7. serie 0001 exclusion preserved
# ---------------------------------------------------------------------------

def test_exclusion_serie_0001_preserved():
    df = pd.DataFrame({
        'numruc': ['20000000001', '20000000002', '20000000003'],
        'serie': ['0001', 'FE01', 'BE01'],
        'documento': ['0001', '0000002', '0000003'],
        'importe_cancelado': [200.0, 100.0, 50.0],
        'banco': ['INTERBANK', 'INTERBANK', 'INTERBANK'],
        'fecha_cancelacion': [_DATE, _DATE, _DATE],
    })
    _, excluded, processed = split_ode_rows(df)
    assert len(excluded) == 1
    assert len(processed) == 2
    assert not processed['serie'].str.zfill(4).eq('0001').any()


# ---------------------------------------------------------------------------
# 8. Partial payments / duplicates preserved
# ---------------------------------------------------------------------------

def test_partial_payments_not_merged():
    df = _make_df(
        serie=['FE01', 'FE01'],
        documento=['0000001', '0000001'],
        importe_cancelado=[80.0, 20.0],
    )
    carga = build_carga_dataframe(df, periodo='202608', config=_cfg_serie_map())
    h = carga[carga['DEBE HABER'] == 'H']
    assert len(h) == 2
    assert abs(h['IMPORT_TOTAL'].sum() - 100.0) <= 0.01


# ---------------------------------------------------------------------------
# 9. XLS export — TIP_DOC as TEXT with leading zeros
# ---------------------------------------------------------------------------

def test_xls_tip_doc_text_format(tmp_path):
    carga = build_carga_dataframe(_make_df(), periodo='202608', config=_cfg_serie_map())
    xls_path = tmp_path / 'mixed.xls'
    export_carga_xls(carga, xls_path)

    wb = xlrd.open_workbook(str(xls_path))
    ws = wb.sheet_by_index(0)
    hdrs = [ws.cell_value(0, c) for c in range(ws.ncols)]
    tip_col, dh_col = hdrs.index('TIP_DOC'), hdrs.index('DEBE HABER')

    h_types = []
    for r in range(1, ws.nrows):
        if str(ws.cell_value(r, dh_col)) == 'H':
            cell = ws.cell(r, tip_col)
            assert cell.ctype == xlrd.XL_CELL_TEXT, (
                f'row {r+1}: TIP_DOC must be TEXT, got ctype={cell.ctype}'
            )
            h_types.append(cell.value)

    assert '01' in h_types and '03' in h_types


def test_xls_tip_doc_leading_zeros(tmp_path):
    carga = build_carga_dataframe(_make_df(), periodo='202608', config=_cfg_serie_map())
    xls_path = tmp_path / 'lz.xls'
    export_carga_xls(carga, xls_path)

    wb = xlrd.open_workbook(str(xls_path))
    ws = wb.sheet_by_index(0)
    hdrs = [ws.cell_value(0, c) for c in range(ws.ncols)]
    tip_col, dh_col = hdrs.index('TIP_DOC'), hdrs.index('DEBE HABER')

    for r in range(1, ws.nrows):
        if str(ws.cell_value(r, dh_col)) == 'H':
            val = ws.cell_value(r, tip_col)
            assert val in ('01', '03'), f'row {r+1}: got {val!r}'


# ---------------------------------------------------------------------------
# 10. process_report — synthetic file with explicit tipo_comprobante column
# ---------------------------------------------------------------------------

def _make_synthetic_xlsx(path: Path, with_tipo_col: bool = True) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    cols = ['NumRuc', 'Serie', 'Documento', 'Importe Cancelado', 'Fecha Cancelacion', 'Banco']
    if with_tipo_col:
        cols.append('Tipo Comprobante')
    ws.append(cols)
    rows = [
        ('20111111111', 'FE01', '0000001', 100.0, datetime(2026, 8, 10), 'INTERBANK', '01'),
        ('20222222222', 'BE01', '0000002',  50.0, datetime(2026, 8, 11), 'INTERBANK', '03'),
        ('20333333333', 'FE01', '0000003',  75.0, datetime(2026, 8, 12), 'OTROS',     '01'),
    ]
    for r in rows:
        ws.append(r[:6] + ((r[6],) if with_tipo_col else ()))
    wb.save(str(path))


def test_process_report_explicit_column(tmp_path):
    """Explicit tipo_comprobante column → 01/03 generated correctly."""
    src = tmp_path / 's.xlsx'
    _make_synthetic_xlsx(src, with_tipo_col=True)
    result = process_report(src, tmp_path, config={**DEFAULT_CONFIG, 'serie_map': {}},
                            periodo='202608')
    s = result['summary']
    assert s['resultado'] == 'CUADRADO'
    ts = s['tip_doc_summary']
    assert ts['01'] == 2 and ts['03'] == 1 and ts['sin_resolver'] == 0
    assert abs(ts['importe_01'] - 175.0) <= 0.01
    assert abs(ts['importe_03'] -  50.0) <= 0.01


def test_process_report_no_column_empty_map_blocked(tmp_path):
    """No tipo_comprobante column + empty serie_map → generation blocked."""
    src = tmp_path / 's.xlsx'
    _make_synthetic_xlsx(src, with_tipo_col=False)
    with pytest.raises(ValueError):
        process_report(src, tmp_path, config={**DEFAULT_CONFIG, 'serie_map': {}}, periodo='202608')


def test_process_report_no_column_with_serie_map(tmp_path):
    """No tipo_comprobante column + complete serie_map → generation succeeds."""
    src = tmp_path / 's.xlsx'
    _make_synthetic_xlsx(src, with_tipo_col=False)
    result = process_report(src, tmp_path,
                            config={**DEFAULT_CONFIG, 'serie_map': {'FE': '01', 'BE': '03'}},
                            periodo='202608')
    ts = result['summary']['tip_doc_summary']
    assert ts['01'] == 2 and ts['03'] == 1 and ts['sin_resolver'] == 0


# ---------------------------------------------------------------------------
# 11. Accounting balance — mixed multi-bank
# ---------------------------------------------------------------------------

def test_accounting_balance_mixed_multibank():
    df = pd.DataFrame({
        'numruc': ['20000000001', '20000000002', '20000000003'],
        'serie': ['FE01', 'BE01', 'FE01'],
        'documento': ['001', '002', '003'],
        'importe_cancelado': [100.0, 50.0, 75.0],
        'banco': ['INTERBANK', 'BBVA CONTINENTAL', 'OTROS'],
        'fecha_cancelacion': [_DATE, _DATE, _DATE],
    })
    carga = build_carga_dataframe(df, periodo='202608', config=_cfg_serie_map())
    d = carga[carga['DEBE HABER'] == 'D']['IMPORT_TOTAL'].sum()
    h = carga[carga['DEBE HABER'] == 'H']['IMPORT_TOTAL'].sum()
    assert abs(d - h) <= 0.01 and abs(d - 225.0) <= 0.01
    total = next(r for r in build_conciliacion(carga) if r['banco'] == 'TOTAL')
    assert total['estado'] == 'CUADRADO'


# ---------------------------------------------------------------------------
# 12. tip_doc_summary amounts — computed from processed rows, not D-rows
# ---------------------------------------------------------------------------

def test_tip_doc_summary_amounts_correct(tmp_path):
    src = tmp_path / 's.xlsx'
    _make_synthetic_xlsx(src, with_tipo_col=True)
    result = process_report(src, tmp_path, config={**DEFAULT_CONFIG, 'serie_map': {}})
    ts = result['summary']['tip_doc_summary']
    assert abs(ts['importe_01'] - 175.0) <= 0.01   # 100 + 75
    assert abs(ts['importe_03'] -  50.0) <= 0.01
    assert ts['importe_sin_resolver'] == 0.0


# ---------------------------------------------------------------------------
# 13. Validation report — TipoDoc sheet with Cantidad and Importe
# ---------------------------------------------------------------------------

def test_validation_report_tipodoc_has_amounts(tmp_path):
    import openpyxl
    src = tmp_path / 's.xlsx'
    _make_synthetic_xlsx(src, with_tipo_col=True)
    result = process_report(src, tmp_path, config={**DEFAULT_CONFIG, 'serie_map': {}})
    wb = openpyxl.load_workbook(result['report_file'])
    assert 'TipoDoc' in wb.sheetnames
    ws = wb['TipoDoc']
    assert ws.cell(1, 3).value == 'Importe'
    rows = {ws.cell(r, 1).value: (ws.cell(r, 2).value, ws.cell(r, 3).value)
            for r in range(2, ws.max_row + 1)}
    assert rows['Facturas (01)'][0] == 2 and abs(rows['Facturas (01)'][1] - 175.0) <= 0.01
    assert rows['Boletas (03)'][0] == 1 and abs(rows['Boletas (03)'][1] -  50.0) <= 0.01


# ---------------------------------------------------------------------------
# 14. _collect_config — empty combo must not raise IndexError
# ---------------------------------------------------------------------------

def test_collect_config_empty_combo_no_crash():
    """Empty combo.get() must not cause IndexError in _collect_config."""
    pytest.importorskip('tkinter')
    from app import ConversorODEApp
    try:
        app = ConversorODEApp()
    except Exception as exc:
        pytest.skip(f'Tk unavailable: {exc}')
    from tkinter import ttk
    combo = ttk.Combobox(app._inner, values=['01 (Factura)', '03 (Boleta)'], state='readonly')
    app.serie_map_widgets['ZZ99'] = combo  # empty — user hasn't selected
    cfg = app._collect_config()            # must NOT raise IndexError
    assert 'ZZ99' not in cfg.get('serie_map', {}), (
        'Empty combo must not auto-assign ZZ99 to serie_map'
    )
    app.destroy()
