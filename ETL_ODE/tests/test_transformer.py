"""
Regression tests for the ODE accounting ETL pipeline.
Run from the repo root: pytest -q
"""
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
import xlrd

from transformer import (
    DEFAULT_CONFIG,
    OUTPUT_COLUMNS,
    auto_create_bank_config,
    build_carga_dataframe,
    build_conciliacion,
    detect_total_row,
    export_carga_xls,
    get_bank_default_config,
    load_config,
    load_ode_report,
    normalize_amount,
    normalize_document_value,
    normalize_serie,
    parse_date_value,
    process_report,
    read_raw_ode_report,
    serialize_serie_num,
    split_ode_rows,
)

# Tests live in ETL_ODE/tests/; data files are in ETL_ODE/
_ETL_DIR = Path(__file__).resolve().parent.parent
SOURCE_PATH = _ETL_DIR / 'reporte de cancelados agosto 2026.xlsx'
MODEL_PATH = _ETL_DIR / 'FACTURA 07-2026 SHIKINA.xls'


# ---------------------------------------------------------------------------
# 1. Raw split — counts match the UI requirements exactly
# ---------------------------------------------------------------------------

def test_raw_split_counts():
    raw = read_raw_ode_report(SOURCE_PATH)
    raw, excluded, processed = split_ode_rows(raw)
    assert len(raw) == 151, f'expected 151 raw, got {len(raw)}'
    assert len(excluded) == 10, f'expected 10 excluded, got {len(excluded)}'
    assert len(processed) == 141, f'expected 141 processed, got {len(processed)}'


def test_raw_split_total_matches():
    raw = read_raw_ode_report(SOURCE_PATH)
    _, _, processed = split_ode_rows(raw)
    total = round(float(processed['importe_cancelado'].sum()), 2)
    assert total == 26757.00, f'expected S/26,757.00, got {total}'


# ---------------------------------------------------------------------------
# 2. Exclusion
# ---------------------------------------------------------------------------

def test_exclusion_serie_0001():
    df = load_ode_report(SOURCE_PATH)
    assert not df['serie'].str.startswith('0001').any()


# ---------------------------------------------------------------------------
# 3. D-row header fields must be blank
# ---------------------------------------------------------------------------

def test_d_row_document_fields_blank():
    df = load_ode_report(SOURCE_PATH)
    carga = build_carga_dataframe(df, periodo='202608', config={})
    d_rows = carga[carga['DEBE HABER'] == 'D']
    assert d_rows['TIP_ANEXO'].astype(str).str.strip().eq('').all(),  'TIP_ANEXO should be blank on D rows'
    assert d_rows['COD_ANEXO'].astype(str).str.strip().eq('').all(),  'COD_ANEXO should be blank on D rows'
    assert d_rows['TIP_DOC'].astype(str).str.strip().eq('').all(),    'TIP_DOC should be blank on D rows'
    assert d_rows['SERIE_NUM'].astype(str).str.strip().eq('').all(),  'SERIE_NUM should be blank on D rows'
    assert d_rows['FEC_VENC'].astype(str).str.strip().eq('').all(),   'FEC_VENC should be blank on D rows'
    assert d_rows['TIP_CAMBIO'].astype(str).str.strip().eq('').all(), 'TIP_CAMBIO should be blank on D rows'
    assert d_rows['CENTRO_COSTOS'].astype(str).str.strip().eq('').all(), 'CENTRO_COSTOS should be blank on D rows'
    assert d_rows['NRO_FILE'].astype(str).str.strip().eq('').all(),   'NRO_FILE should be blank on D rows'


def test_d_row_fec_doc_is_month_end():
    df = load_ode_report(SOURCE_PATH)
    carga = build_carga_dataframe(df, periodo='202608', config={})
    d_rows = carga[carga['DEBE HABER'] == 'D']
    for fec in d_rows['FEC_DOC']:
        dt = pd.to_datetime(fec, errors='coerce')
        assert dt.day == 31 and dt.month == 8 and dt.year == 2026, f'D-row FEC_DOC must be 31/08/2026, got {fec}'


# ---------------------------------------------------------------------------
# 4. Conciliation — exact expected values
# ---------------------------------------------------------------------------

EXPECTED_CONC = [
    ('INTERBANK',        '0001', 79,  20002.69, 20002.69, 0.00),
    ('BBVA CONTINENTAL', '0002', 15,  1720.00,  1720.00,  0.00),
    ('DE LA NACION',     '0003', 22,  4942.00,  4942.00,  0.00),
    ('OTROS',            '0004', 25,    92.31,    92.31,  0.00),
    ('TOTAL',           'TOTAL', 141, 26757.00, 26757.00, 0.00),
]


def test_conciliacion_exact_values():
    df = load_ode_report(SOURCE_PATH)
    carga = build_carga_dataframe(df, periodo='202608', config={})
    conc = build_conciliacion(carga)
    assert len(conc) == 5, f'expected 5 conciliation rows (4 banks + total), got {len(conc)}'
    for i, (banco, comp, cantidad, debito, credito, diferencia) in enumerate(EXPECTED_CONC):
        row = conc[i]
        assert row['banco'] == banco, f'row {i} banco: expected {banco!r}, got {row["banco"]!r}'
        assert row['comprobante'] == comp, f'row {i} comprobante: expected {comp!r}, got {row["comprobante"]!r}'
        assert row['cantidad'] == cantidad, f'row {i} cantidad: expected {cantidad}, got {row["cantidad"]}'
        assert abs(row['debito'] - debito) <= 0.01, f'row {i} debito: expected {debito}, got {row["debito"]}'
        assert abs(row['credito'] - credito) <= 0.01, f'row {i} credito: expected {credito}, got {row["credito"]}'
        assert abs(row['diferencia'] - diferencia) <= 0.01, f'row {i} diferencia: expected {diferencia}, got {row["diferencia"]}'
        assert row['estado'] == 'CUADRADO', f'row {i} estado should be CUADRADO'


# ---------------------------------------------------------------------------
# 5. Output workbook: exactly 22 columns, no _banco, formats applied
# ---------------------------------------------------------------------------

def test_output_has_exactly_22_columns(tmp_path):
    result = process_report(
        SOURCE_PATH, tmp_path, config={},
        fecha_calculo='fecha_cancelacion', periodo='202608',
        model_path=MODEL_PATH,
    )
    # Primary output is .xls — verify via xlrd
    wb = xlrd.open_workbook(result['output_file'])
    ws = wb.sheet_by_index(0)
    headers = [ws.cell_value(0, c) for c in range(ws.ncols)]
    assert headers == list(OUTPUT_COLUMNS), f'column mismatch: {headers}'
    assert '_banco' not in headers, '_banco aux column must not appear in exported file'


def test_output_date_and_number_formats(tmp_path):
    result = process_report(
        SOURCE_PATH, tmp_path, config={},
        fecha_calculo='fecha_cancelacion', periodo='202608',
    )
    # Backup .xlsx retains openpyxl number format metadata for review
    wb = openpyxl.load_workbook(result['output_file_xlsx'])
    ws = wb.active
    # Check first data row (row 2); D-row FEC_DOC is col E, IMPORT_TOTAL is col L
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            if cell.column_letter in ('E', 'J', 'N') and cell.value not in ('', None):
                assert 'D' in cell.number_format.upper() or 'Y' in cell.number_format.upper(), \
                    f'Expected date format on {cell.coordinate}, got {cell.number_format!r}'
            if cell.column_letter == 'L' and cell.value not in ('', None):
                assert '0' in cell.number_format, \
                    f'Expected numeric format on {cell.coordinate}, got {cell.number_format!r}'
        break  # one row is enough to verify formats were applied


# ---------------------------------------------------------------------------
# 6. Full process_report end-to-end
# ---------------------------------------------------------------------------

def test_process_report_real_report_generates_valid_carga(tmp_path):
    result = process_report(
        SOURCE_PATH, tmp_path, config={},
        fecha_calculo='fecha_cancelacion', periodo='202608',
        model_path=MODEL_PATH,
    )
    s = result['summary']
    assert s['cantidad_original'] == 151
    assert s['cantidad_excluida'] == 10
    assert s['cantidad_procesada'] == 141
    assert abs(s['total_procesado'] - 26757.00) <= 0.01
    assert abs(s['total_debito'] - 26757.00) <= 0.01
    assert abs(s['total_credito'] - 26757.00) <= 0.01
    assert s['diferencia'] == 0.00
    assert s['resultado'] == 'CUADRADO'

    carga = result['output_df']
    # 141 H rows + 4 D rows = 145
    assert len(carga) == 145, f'expected 145 output rows, got {len(carga)}'
    # Primary output is .xls
    assert Path(result['output_file']).exists()
    assert Path(result['output_file']).suffix == '.xls', 'primary output must be .xls'
    # Backup .xlsx also generated
    assert Path(result['output_file_xlsx']).exists()
    assert Path(result['report_file']).exists()


# ---------------------------------------------------------------------------
# 7. analyze_source helper mirrors UI counts
# ---------------------------------------------------------------------------

def test_analyze_source_counts():
    from app import analyze_source
    raw, excluded, processed = analyze_source(str(SOURCE_PATH))
    assert len(raw) == 151
    assert len(excluded) == 10
    assert len(processed) == 141
    assert abs(float(processed['importe_cancelado'].sum()) - 26757.00) <= 0.01


# ---------------------------------------------------------------------------
# 8. Existing unit tests
# ---------------------------------------------------------------------------

def test_concatenacion_serie_num():
    assert serialize_serie_num({'serie': 'FE01', 'documento': '0029546'}) == 'FE010029546'


def test_totales_por_banco():
    df = load_ode_report(SOURCE_PATH)
    totals = df.groupby('banco')['importe_cancelado'].sum()
    for b in ('INTERBANK', 'BBVA CONTINENTAL', 'DE LA NACION', 'OTROS'):
        assert b in totals.index


def test_lectura_fechas_como_texto():
    dt = parse_date_value('25/08/2026')
    assert dt is not None and dt.day == 25 and dt.month == 8 and dt.year == 2026


def test_lectura_importes_con_comas():
    assert abs(normalize_amount('2,50') - 2.5) < 1e-9
    assert abs(normalize_amount('1.234,56') - 1234.56) < 1e-9
    assert abs(normalize_amount('1,234.56') - 1234.56) < 1e-9


def test_preservacion_ruc_y_ceros():
    assert normalize_document_value('02012345678') == '02012345678'
    assert normalize_document_value('000045') == '000045'
    assert normalize_serie('FE01') == 'FE01'


def test_deteccion_fila_de_totales():
    assert detect_total_row('TOTAL') is True
    assert detect_total_row('Reporte de Comprobantes Cancelados') is False


def test_creacion_automatica_banco_desconocido():
    cfg = {'bancos': {}}
    bank = auto_create_bank_config(cfg, 'BANCO NUEVO')
    assert bank['cuenta_contable'] == '10100011'
    assert bank['medio_pago'] == '008'


def test_otros_usa_cuenta_caja_y_medio_vacio():
    assert DEFAULT_CONFIG['bancos']['OTROS'] == {'cuenta_contable': '10100001', 'medio_pago': ''}
    assert get_bank_default_config('OTROS') == {'cuenta_contable': '10100001', 'medio_pago': ''}
    df = load_ode_report(SOURCE_PATH)
    out = build_carga_dataframe(df, periodo='202608', config={})

    otros_d = out[(out['_banco'] == 'OTROS') & (out['DEBE HABER'] == 'D')]
    assert not otros_d.empty, 'No OTROS D-rows found'
    assert otros_d.iloc[0]['CTA_CONTABLE'] == '10100001'
    assert otros_d.iloc[0]['MEDIO_PAGO'] == ''

    otros_h = out[(out['_banco'] == 'OTROS') & (out['DEBE HABER'] == 'H')]
    assert not otros_h.empty, 'No OTROS H-rows found'
    for idx, row in otros_h.iterrows():
        assert row['MEDIO_PAGO'] == '', (
            f'H-row at index {idx} has non-empty MEDIO_PAGO: {row["MEDIO_PAGO"]!r}'
        )


def test_otros_excel_medio_pago_realmente_vacio(tmp_path):
    """MEDIO_PAGO must be genuinely empty in the exported .xls for all OTROS rows."""
    df = load_ode_report(SOURCE_PATH)
    carga = build_carga_dataframe(df, periodo='202608', config={})

    # All OTROS rows (D and H) must have empty MEDIO_PAGO in the DataFrame
    otros_rows = carga[carga['_banco'] == 'OTROS']
    assert not otros_rows.empty, 'No OTROS rows found in carga DataFrame'
    for idx, row in otros_rows.iterrows():
        assert row['MEDIO_PAGO'] == '', (
            f'DataFrame OTROS row {idx} (DEBE HABER={row["DEBE HABER"]!r}) '
            f'has non-empty MEDIO_PAGO: {row["MEDIO_PAGO"]!r}'
        )

    otros_d = otros_rows[otros_rows['DEBE HABER'] == 'D']
    assert not otros_d.empty, 'No OTROS D-rows found'
    otros_comprobante = str(otros_d.iloc[0]['COMPROBANTE'])

    result = process_report(
        SOURCE_PATH, tmp_path, config={},
        fecha_calculo='fecha_cancelacion', periodo='202608',
    )
    # Verify in the primary .xls via xlrd
    wb = xlrd.open_workbook(result['output_file'])
    ws = wb.sheet_by_index(0)
    headers = [ws.cell_value(0, c) for c in range(ws.ncols)]
    medio_col = headers.index('MEDIO_PAGO')
    cta_col   = headers.index('CTA_CONTABLE')
    dh_col    = headers.index('DEBE HABER')
    comp_col  = headers.index('COMPROBANTE')

    otros_xls_rows = 0
    for row_idx in range(1, ws.nrows):
        cta  = str(ws.cell_value(row_idx, cta_col))
        dh   = str(ws.cell_value(row_idx, dh_col))
        comp_cell = ws.cell(row_idx, comp_col)
        comp = str(comp_cell.value) if comp_cell.ctype != xlrd.XL_CELL_EMPTY else ''
        # strip trailing '.0' that xlrd may add for numeric-formatted text
        if comp.endswith('.0'):
            comp = comp[:-2]

        is_otros_d = (cta == '10100001' and dh == 'D')
        is_otros_h = (comp == otros_comprobante and dh == 'H')
        if is_otros_d or is_otros_h:
            otros_xls_rows += 1
            medio_val = ws.cell_value(row_idx, medio_col)
            assert medio_val in (None, ''), (
                f'xls row {row_idx + 1}: MEDIO_PAGO should be empty, got {medio_val!r}'
            )

    assert otros_xls_rows > 0, 'No OTROS rows identified in the exported .xls'


def test_config_migracion_legacy_otros(tmp_path):
    cfg_path = tmp_path / 'config.json'
    cfg_path.write_text('{\n  "bancos": {\n    "OTROS": {\n      "cuenta_contable": "10100011",\n      "medio_pago": "008"\n    }\n  }\n}', encoding='utf-8')
    loaded = load_config(cfg_path)
    assert loaded['bancos']['OTROS'] == {'cuenta_contable': '10100001', 'medio_pago': ''}

    custom_path = tmp_path / 'custom.json'
    custom_path.write_text('{\n  "bancos": {\n    "OTROS": {\n      "cuenta_contable": "99000001",\n      "medio_pago": "007"\n    }\n  }\n}', encoding='utf-8')
    custom_loaded = load_config(custom_path)
    assert custom_loaded['bancos']['OTROS'] == {'cuenta_contable': '99000001', 'medio_pago': '007'}


def test_exportacion_sin_nan_ni_nat():
    df = load_ode_report(SOURCE_PATH)
    out = build_carga_dataframe(df, periodo='202608', config={})
    assert not out[OUTPUT_COLUMNS].isnull().values.any()
    text = out[OUTPUT_COLUMNS].astype(str)
    assert not text.apply(lambda col: col.str.lower().str.contains('nan')).any().any()
    assert not text.apply(lambda col: col.str.lower().str.contains('nat')).any().any()


def test_model_file_uses_expected_structure():
    model = pd.read_excel(MODEL_PATH, header=0)
    assert list(model.columns) == list(OUTPUT_COLUMNS)


def test_igualdad_debe_haber():
    df = load_ode_report(SOURCE_PATH)
    out = build_carga_dataframe(df, periodo='202608', config={})
    debito = out[out['DEBE HABER'] == 'D']['IMPORT_TOTAL'].sum()
    credito = out[out['DEBE HABER'] == 'H']['IMPORT_TOTAL'].sum()
    assert abs(debito - credito) <= 0.01


def test_ui_allows_empty_medio_pago_and_pending_state(tmp_path):
    from app import ConversorODEApp
    import tkinter as tk

    app = ConversorODEApp()
    app.bank_entries['OTROS'] = {'cuenta': tk.Entry(app._inner), 'medio': tk.Entry(app._inner)}
    app.bank_entries['OTROS']['cuenta'].insert(0, '10100001')
    app.bank_entries['OTROS']['medio'].insert(0, '')
    app.config = app._collect_config()
    cfg_path = tmp_path / 'config.json'
    from transformer import save_config
    save_config(cfg_path, app.config)
    loaded = load_config(cfg_path)
    assert loaded['bancos']['OTROS']['medio_pago'] == ''

    raw, excluded, processed = split_ode_rows(read_raw_ode_report(SOURCE_PATH))
    app.file_label.set(str(SOURCE_PATH))

    # --- Escenario con errores: debe mostrar REVISAR (no LISTO PARA GENERAR) ---
    errors = ['Error de prueba: encabezado no encontrado']
    app._refresh_analysis(raw, excluded, processed, errors, [], [], [])
    assert app.cards['balance'].cget('text') == 'PENDIENTE', \
        'balance debe ser PENDIENTE cuando hay errores'
    assert app.cards['status'].cget('text') == 'REVISAR', \
        f"cards['status'] debe ser REVISAR, fue {app.cards['status'].cget('text')!r}"
    assert app.status_label.cget('text') == 'REVISAR', \
        f"status_label debe ser REVISAR, fue {app.status_label.cget('text')!r}"

    # --- Escenario sin errores: debe mostrar LISTO PARA GENERAR ---
    app._refresh_analysis(raw, excluded, processed, [], [], ['Archivo legible'], [])
    assert app.cards['balance'].cget('text') == 'PENDIENTE'
    assert app.cards['status'].cget('text') == 'LISTO PARA GENERAR'
    assert app.status_label.cget('text') == 'LISTO PARA GENERAR'

    from tkinter import messagebox
    original_showinfo = messagebox.showinfo
    messagebox.showinfo = lambda *args, **kwargs: None
    try:
        app._finish_generation(process_report(SOURCE_PATH, tmp_path, config={}, fecha_calculo='fecha_cancelacion', periodo='202608'))
    finally:
        messagebox.showinfo = original_showinfo
    assert app.cards['balance'].cget('text') == 'S/0.00'
    assert app.cards['status'].cget('text') == 'CUADRADO'
    assert app.status_label.cget('text') == 'CUADRADO'
    app.destroy()


# ---------------------------------------------------------------------------
# 9. GUI smoke test — window builds without error and action buttons exist
# ---------------------------------------------------------------------------

def test_gui_smoke_and_buttons():
    pytest.importorskip('tkinter')
    from app import ConversorODEApp
    try:
        app = ConversorODEApp()
    except Exception as exc:
        pytest.skip(f'Tk unavailable in this environment: {exc}')
    app.update_idletasks()
    expected_buttons = {
        'Analizar archivos',
        'Generar carga',
        'Abrir carpeta de salida',
        'Restablecer',
        'Guardar configuración',
    }
    assert expected_buttons == set(app.action_buttons.keys()), \
        f'Missing buttons: {expected_buttons - set(app.action_buttons.keys())}'
    assert hasattr(app, '_canvas'), 'Canvas scroll area not created'
    assert hasattr(app, '_inner'), 'Inner scrollable frame not created'
    app.destroy()


# ---------------------------------------------------------------------------
# 10. XLS output — mandatory checks for macro compatibility
# ---------------------------------------------------------------------------

def _xls_workbook_and_headers(result: dict):
    """Helper: open the primary .xls and return (wb, ws, headers_list)."""
    wb = xlrd.open_workbook(result['output_file'])
    ws = wb.sheet_by_index(0)
    headers = [ws.cell_value(0, c) for c in range(ws.ncols)]
    return wb, ws, headers


def test_xls_file_exists(tmp_path):
    """1. Primary .xls file is present on disk."""
    result = process_report(
        SOURCE_PATH, tmp_path, config={},
        fecha_calculo='fecha_cancelacion', periodo='202608',
    )
    xls_path = Path(result['output_file'])
    assert xls_path.exists(), f'.xls not found: {xls_path}'
    assert xls_path.suffix == '.xls', f'expected .xls extension, got {xls_path.suffix!r}'


def test_xls_readable_with_xlrd(tmp_path):
    """2. xlrd can open the file without error (proves it is real BIFF8, not a renamed .xlsx)."""
    result = process_report(
        SOURCE_PATH, tmp_path, config={},
        fecha_calculo='fecha_cancelacion', periodo='202608',
    )
    wb = xlrd.open_workbook(result['output_file'])
    assert wb.nsheets >= 1
    ws = wb.sheet_by_index(0)
    assert ws.nrows >= 2, 'xls must have header + at least one data row'


def test_xls_has_exactly_22_columns(tmp_path):
    """3. Exactly 22 columns matching OUTPUT_COLUMNS; no _banco."""
    result = process_report(
        SOURCE_PATH, tmp_path, config={},
        fecha_calculo='fecha_cancelacion', periodo='202608',
    )
    _, ws, headers = _xls_workbook_and_headers(result)
    assert len(headers) == 22, f'expected 22 columns, got {len(headers)}: {headers}'
    assert headers == list(OUTPUT_COLUMNS), f'column order/names mismatch: {headers}'
    assert '_banco' not in headers, '_banco must not appear in exported .xls'


def test_xls_has_145_data_rows(tmp_path):
    """4. 145 data rows (header row excluded)."""
    result = process_report(
        SOURCE_PATH, tmp_path, config={},
        fecha_calculo='fecha_cancelacion', periodo='202608',
    )
    _, ws, _ = _xls_workbook_and_headers(result)
    data_rows = ws.nrows - 1  # subtract header
    assert data_rows == 145, f'expected 145 data rows, got {data_rows}'


def test_xls_has_4d_141h_rows(tmp_path):
    """5. Exactly 4 D-rows and 141 H-rows."""
    result = process_report(
        SOURCE_PATH, tmp_path, config={},
        fecha_calculo='fecha_cancelacion', periodo='202608',
    )
    _, ws, headers = _xls_workbook_and_headers(result)
    dh_col = headers.index('DEBE HABER')
    d_count = h_count = 0
    for row_idx in range(1, ws.nrows):
        val = str(ws.cell_value(row_idx, dh_col))
        if val == 'D':
            d_count += 1
        elif val == 'H':
            h_count += 1
    assert d_count == 4,   f'expected 4 D-rows, got {d_count}'
    assert h_count == 141, f'expected 141 H-rows, got {h_count}'


def test_xls_totals_26757(tmp_path):
    """6. Debe and Haber each sum to S/26,757.00; difference = S/0.00."""
    result = process_report(
        SOURCE_PATH, tmp_path, config={},
        fecha_calculo='fecha_cancelacion', periodo='202608',
    )
    _, ws, headers = _xls_workbook_and_headers(result)
    dh_col     = headers.index('DEBE HABER')
    import_col = headers.index('IMPORT_TOTAL')
    total_d = total_h = 0.0
    for row_idx in range(1, ws.nrows):
        dh  = str(ws.cell_value(row_idx, dh_col))
        amt = ws.cell_value(row_idx, import_col)
        if isinstance(amt, (int, float)):
            if dh == 'D':
                total_d += amt
            elif dh == 'H':
                total_h += amt
    assert abs(round(total_d, 2) - 26757.00) <= 0.01, f'D total: expected 26757.00, got {total_d:.2f}'
    assert abs(round(total_h, 2) - 26757.00) <= 0.01, f'H total: expected 26757.00, got {total_h:.2f}'
    assert abs(total_d - total_h) <= 0.01, f'difference must be 0, got {abs(total_d - total_h):.2f}'


def test_xls_otros_uses_cuenta_10100001(tmp_path):
    """7. All D-rows for OTROS use account 10100001."""
    result = process_report(
        SOURCE_PATH, tmp_path, config={},
        fecha_calculo='fecha_cancelacion', periodo='202608',
    )
    _, ws, headers = _xls_workbook_and_headers(result)
    cta_col = headers.index('CTA_CONTABLE')
    dh_col  = headers.index('DEBE HABER')

    otros_d_rows = []
    for row_idx in range(1, ws.nrows):
        cta = str(ws.cell_value(row_idx, cta_col))
        dh  = str(ws.cell_value(row_idx, dh_col))
        if cta == '10100001' and dh == 'D':
            otros_d_rows.append(row_idx)

    assert len(otros_d_rows) > 0, 'No OTROS D-rows (account 10100001) found in .xls'
    for row_idx in otros_d_rows:
        cta = str(ws.cell_value(row_idx, cta_col))
        assert cta == '10100001', f'row {row_idx + 1}: CTA_CONTABLE expected 10100001, got {cta!r}'


def test_xls_otros_medio_pago_empty(tmp_path):
    """8. MEDIO_PAGO is empty for all D and H rows belonging to OTROS."""
    df = load_ode_report(SOURCE_PATH)
    carga = build_carga_dataframe(df, periodo='202608', config={})
    otros_d = carga[(carga['_banco'] == 'OTROS') & (carga['DEBE HABER'] == 'D')]
    assert not otros_d.empty, 'No OTROS D-rows in DataFrame'
    otros_comprobante = str(otros_d.iloc[0]['COMPROBANTE'])

    result = process_report(
        SOURCE_PATH, tmp_path, config={},
        fecha_calculo='fecha_cancelacion', periodo='202608',
    )
    _, ws, headers = _xls_workbook_and_headers(result)
    cta_col   = headers.index('CTA_CONTABLE')
    dh_col    = headers.index('DEBE HABER')
    comp_col  = headers.index('COMPROBANTE')
    medio_col = headers.index('MEDIO_PAGO')

    checked = 0
    for row_idx in range(1, ws.nrows):
        cta  = str(ws.cell_value(row_idx, cta_col))
        dh   = str(ws.cell_value(row_idx, dh_col))
        comp_raw = ws.cell_value(row_idx, comp_col)
        comp = str(comp_raw).rstrip('.0') if str(comp_raw).endswith('.0') else str(comp_raw)

        is_otros_d = (cta == '10100001' and dh == 'D')
        is_otros_h = (comp == otros_comprobante and dh == 'H')
        if is_otros_d or is_otros_h:
            checked += 1
            medio_val = ws.cell_value(row_idx, medio_col)
            assert medio_val in (None, ''), (
                f'xls row {row_idx + 1}: MEDIO_PAGO must be empty for OTROS, got {medio_val!r}'
            )

    assert checked > 0, 'No OTROS rows found in .xls to verify MEDIO_PAGO'


def test_xls_no_banco_column(tmp_path):
    """9. The string _banco must not appear anywhere in the .xls headers."""
    result = process_report(
        SOURCE_PATH, tmp_path, config={},
        fecha_calculo='fecha_cancelacion', periodo='202608',
    )
    _, _, headers = _xls_workbook_and_headers(result)
    assert '_banco' not in headers, f'_banco aux column must not be exported; headers: {headers}'
    assert all('_banco' not in str(h) for h in headers), \
        f'Unexpected _banco variant in headers: {headers}'


def test_validation_report_is_xlsx(tmp_path):
    """10. REPORTE_VALIDACION always generates as .xlsx regardless of .xls change."""
    result = process_report(
        SOURCE_PATH, tmp_path, config={},
        fecha_calculo='fecha_cancelacion', periodo='202608',
    )
    report_path = Path(result['report_file'])
    assert report_path.exists(), f'Validation report not found: {report_path}'
    assert report_path.suffix == '.xlsx', f'expected .xlsx report, got {report_path.suffix!r}'
    # Verify it is a valid xlsx (openpyxl can open it)
    wb = openpyxl.load_workbook(report_path)
    assert 'Resumen' in wb.sheetnames, 'Resumen sheet missing from validation report'


# ---------------------------------------------------------------------------
# 11. XLS cell-type compatibility with model (FACTURA 07-2026 SHIKINA.xls)
# ---------------------------------------------------------------------------

def test_xls_cell_types_match_model(tmp_path):
    """Compare cell types of generated .xls against the original model file.

    For each OUTPUT_COLUMN the test verifies that non-empty cells in both files
    carry the same xlrd ctype.  A representative D-row and H-row are used.
    Then specific mandatory assertions enforce ANO_MES / DOC_ANULADO are NUMBER,
    key text fields are TEXT, dates are DATE, IMPORT_TOTAL is NUMBER, and the
    sheet is named exactly 'Hoja1'.
    """
    result = process_report(
        SOURCE_PATH, tmp_path, config={},
        fecha_calculo='fecha_cancelacion', periodo='202608',
        model_path=MODEL_PATH,
    )

    gen_wb  = xlrd.open_workbook(result['output_file'])
    gen_ws  = gen_wb.sheet_by_index(0)
    mod_wb  = xlrd.open_workbook(str(MODEL_PATH))
    mod_ws  = mod_wb.sheet_by_index(0)

    # Sheet name
    assert gen_ws.name == 'Hoja1', f'sheet name must be Hoja1, got {gen_ws.name!r}'

    # Header order must match model exactly
    gen_headers = [gen_ws.cell_value(0, c) for c in range(gen_ws.ncols)]
    mod_headers = [mod_ws.cell_value(0, c) for c in range(mod_ws.ncols)]
    assert gen_headers == mod_headers, (
        f'Header mismatch.\n  generated: {gen_headers}\n  model:     {mod_headers}'
    )

    # Locate a D-row and an H-row in the generated file
    dh_col = gen_headers.index('DEBE HABER')
    gen_d_row = gen_h_row = None
    for r in range(1, gen_ws.nrows):
        dh = str(gen_ws.cell_value(r, dh_col))
        if gen_d_row is None and dh == 'D':
            gen_d_row = r
        if gen_h_row is None and dh == 'H':
            gen_h_row = r
        if gen_d_row is not None and gen_h_row is not None:
            break

    assert gen_d_row is not None, 'No D-row found in generated .xls'
    assert gen_h_row is not None, 'No H-row found in generated .xls'

    # Locate matching rows in model
    mod_d_row = mod_h_row = None
    for r in range(1, mod_ws.nrows):
        dh = str(mod_ws.cell_value(r, dh_col))
        if mod_d_row is None and dh == 'D':
            mod_d_row = r
        if mod_h_row is None and dh == 'H':
            mod_h_row = r
        if mod_d_row is not None and mod_h_row is not None:
            break

    assert mod_d_row is not None, 'No D-row in model .xls'
    assert mod_h_row is not None, 'No H-row in model .xls'

    # For each column: when BOTH files have a non-empty cell in the same row type,
    # their ctypes must match.
    EMPTY_CTYPES = {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}
    for col_idx, col_name in enumerate(gen_headers):
        for label, gen_row, mod_row in [('D', gen_d_row, mod_d_row),
                                         ('H', gen_h_row, mod_h_row)]:
            gen_cell = gen_ws.cell(gen_row, col_idx)
            mod_cell = mod_ws.cell(mod_row, col_idx)
            # Only compare when the model cell has actual content
            if mod_cell.ctype in EMPTY_CTYPES:
                continue
            assert gen_cell.ctype == mod_cell.ctype, (
                f'{label}-row col [{col_idx}] {col_name!r}: '
                f'expected ctype={mod_cell.ctype}, got {gen_cell.ctype} '
                f'(gen_val={gen_cell.value!r}, mod_val={mod_cell.value!r})'
            )

    # Mandatory type assertions (independent of model row lookup)
    ano_col  = gen_headers.index('ANO_MES')
    doc_col  = gen_headers.index('DOC_ANULADO')
    cta_col  = gen_headers.index('CTA_CONTABLE')
    sub_col  = gen_headers.index('SUB_DIARIO')
    comp_col = gen_headers.index('COMPROBANTE')
    cod_col  = gen_headers.index('COD_ANEXO')
    ser_col  = gen_headers.index('SERIE_NUM')
    imp_col  = gen_headers.index('IMPORT_TOTAL')
    fec_col  = gen_headers.index('FEC_DOC')

    # ANO_MES and DOC_ANULADO must be NUMBER in every data row
    for r in range(1, gen_ws.nrows):
        assert gen_ws.cell(r, ano_col).ctype == xlrd.XL_CELL_NUMBER, (
            f'Row {r + 1}: ANO_MES ctype must be NUMBER, got {gen_ws.cell(r, ano_col).ctype}'
        )
        assert gen_ws.cell(r, doc_col).ctype == xlrd.XL_CELL_NUMBER, (
            f'Row {r + 1}: DOC_ANULADO ctype must be NUMBER, got {gen_ws.cell(r, doc_col).ctype}'
        )
        assert gen_ws.cell(r, imp_col).ctype == xlrd.XL_CELL_NUMBER, (
            f'Row {r + 1}: IMPORT_TOTAL ctype must be NUMBER, got {gen_ws.cell(r, imp_col).ctype}'
        )

    # Text columns must be TEXT when they have content
    text_checks = [
        (cta_col,  'CTA_CONTABLE'),
        (sub_col,  'SUB_DIARIO'),
        (comp_col, 'COMPROBANTE'),
    ]
    for col_idx, col_name in text_checks:
        for r in range(1, gen_ws.nrows):
            cell = gen_ws.cell(r, col_idx)
            if cell.ctype not in EMPTY_CTYPES and cell.value != '':
                assert cell.ctype == xlrd.XL_CELL_TEXT, (
                    f'Row {r + 1}: {col_name} must be TEXT, got ctype={cell.ctype} val={cell.value!r}'
                )

    # COD_ANEXO and SERIE_NUM are TEXT on H-rows (which always have content)
    for r in range(1, gen_ws.nrows):
        dh = str(gen_ws.cell_value(r, dh_col))
        if dh == 'H':
            assert gen_ws.cell(r, cod_col).ctype == xlrd.XL_CELL_TEXT, (
                f'Row {r + 1} (H): COD_ANEXO must be TEXT, got {gen_ws.cell(r, cod_col).ctype}'
            )
            assert gen_ws.cell(r, ser_col).ctype == xlrd.XL_CELL_TEXT, (
                f'Row {r + 1} (H): SERIE_NUM must be TEXT, got {gen_ws.cell(r, ser_col).ctype}'
            )

    # FEC_DOC is DATE on every row
    for r in range(1, gen_ws.nrows):
        assert gen_ws.cell(r, fec_col).ctype == xlrd.XL_CELL_DATE, (
            f'Row {r + 1}: FEC_DOC must be DATE, got {gen_ws.cell(r, fec_col).ctype}'
        )
