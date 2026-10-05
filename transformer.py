from __future__ import annotations

import json
import logging
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

import numpy as np
import pandas as pd

OUTPUT_COLUMNS = [
    'CTA_CONTABLE', 'ANO_MES', 'SUB_DIARIO', 'COMPROBANTE', 'FEC_DOC', 'TIP_ANEXO',
    'COD_ANEXO', 'TIP_DOC', 'SERIE_NUM', 'FEC_VENC', 'MONEDA', 'IMPORT_TOTAL',
    'TIP_CONEVR.', 'FEC_REG', 'TIP_CAMBIO', 'GLOSA', 'CENTRO_COSTOS', 'GLOSA_MOV',
    'DOC_ANULADO', 'DEBE HABER', 'MEDIO_PAGO', 'NRO_FILE'
]

BANK_PRIORITY = ['INTERBANK', 'BBVA CONTINENTAL', 'DE LA NACION', 'OTROS']
DEFAULT_CONFIG = {
    'cuenta_clientes': '12120001',
    'sub_diario': '01',
    'tipo_anexo': '02',
    'tipo_documento': '01',
    'moneda': 'MN',
    'tipo_conversion': 'VTA',
    'doc_anulado': '0',
    'glosa': 'COBRANZA DEL MES',
    'fecha_documento': 'fecha_cancelacion',
    'bancos': {
        'INTERBANK': {'cuenta_contable': '10100011', 'medio_pago': '008'},
        'BBVA CONTINENTAL': {'cuenta_contable': '10100011', 'medio_pago': '008'},
        'DE LA NACION': {'cuenta_contable': '10100011', 'medio_pago': '008'},
        'OTROS': {'cuenta_contable': '10100011', 'medio_pago': '008'},
    },
}


def get_app_dir() -> Path:
    if getattr(sys, 'frozen', False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


LOGGER = logging.getLogger('conversor_ode')
if not LOGGER.handlers:
    log_dir = get_app_dir() / 'logs'
    log_dir.mkdir(parents=True, exist_ok=True)
    LOGGER.setLevel(logging.INFO)
    handler = logging.FileHandler(log_dir / 'conversor_ode.log', encoding='utf-8')
    formatter = logging.Formatter('%(asctime)s | %(levelname)s | %(message)s')
    handler.setFormatter(formatter)
    LOGGER.addHandler(handler)


def normalize_text(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ''
    if isinstance(value, (pd.Timestamp, datetime, date)):
        return value.strftime('%d/%m/%Y')
    if isinstance(value, (float, np.floating)):
        if np.isnan(value):
            return ''
        if value.is_integer():
            return str(int(value))
        return str(value)
    text = str(value).strip()
    return text


def normalize_header_name(value: Any) -> str:
    cleaned = normalize_text(value)
    cleaned = unicodedata.normalize('NFKD', cleaned).encode('ascii', 'ignore').decode('utf-8')
    cleaned = cleaned.lower().replace(' ', '_').replace('-', '_').replace('.', '').replace('/', '_')
    cleaned = re.sub(r'_+', '_', cleaned).strip('_')
    return cleaned


def detect_total_row(value: Any) -> bool:
    text = normalize_text(value).upper()
    return text.startswith('TOTAL') or 'TOTAL' in text


def _remove_accents(value: str) -> str:
    return unicodedata.normalize('NFKD', value).encode('ascii', 'ignore').decode('utf-8')


def normalize_bank_name(value: Any) -> str:
    text = normalize_text(value).upper().strip()
    if not text:
        return ''
    text = _remove_accents(text)
    if 'INTERBANK' in text:
        return 'INTERBANK'
    if 'BBVA' in text or 'CONTINENTAL' in text:
        return 'BBVA CONTINENTAL'
    if 'NACION' in text:
        return 'DE LA NACION'
    if 'OTROS' in text:
        return 'OTROS'
    return text.title()


def normalize_serie(value: Any) -> str:
    text = normalize_text(value).upper().strip()
    return _remove_accents(text)


def normalize_document_value(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ''
    text = normalize_text(value)
    if not text:
        return ''
    if isinstance(value, (float, np.floating)):
        return format(float(value), 'f').rstrip('0').rstrip('.')
    if 'E-' in text.upper() or 'E+' in text.upper():
        try:
            return format(float(text), 'f').rstrip('0').rstrip('.')
        except Exception:
            pass
    return text


def serialize_serie_num(row: Dict[str, Any]) -> str:
    serie = normalize_serie(row.get('serie', ''))
    documento = normalize_document_value(row.get('documento', ''))
    if not serie and not documento:
        return ''
    if not documento:
        return serie
    if not serie:
        return documento
    return f'{serie}{documento}'


def normalize_amount(value: Any) -> float:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return 0.0
    if isinstance(value, (int, float, np.integer, np.floating)):
        return float(value)
    text = normalize_text(value).replace('\u00a0', '').replace(' ', '')
    if not text or text.lower() in {'nan', 'nat', 'none', ''}:
        return 0.0
    if text.startswith('(') and text.endswith(')'):
        text = f'-{text[1:-1]}'
    text = text.replace('S/', '').replace('$', '')
    if '.' in text and ',' in text:
        if text.rfind(',') > text.rfind('.'):
            text = text.replace('.', '').replace(',', '.')
        else:
            text = text.replace(',', '')
    elif ',' in text:
        if text.count(',') > 1:
            text = text.replace(',', '')
        else:
            left, right = text.split(',', 1)
            if len(right) in (2, 3):
                if len(right) == 3 and left.isdigit() and int(left) > 999:
                    text = f'{left}{right}'
                else:
                    text = f'{left}.{right}'
            else:
                text = f'{left}.{right}'
    elif '.' in text and text.count('.') > 1:
        parts = text.split('.')
        if len(parts[-1]) == 3:
            text = ''.join(parts)
    try:
        return float(text)
    except Exception:
        return 0.0


def parse_date_value(value: Any) -> Optional[datetime]:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, (int, float, np.integer, np.floating)):
        try:
            return pd.to_datetime(value, unit='D', origin='1899-12-30', errors='coerce').to_pydatetime()
        except Exception:
            return None
    text = normalize_text(value)
    if not text:
        return None
    text = text.replace(' 00:00:00', '').replace('T', ' ')
    for fmt in ['%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d', '%Y-%m-%d', '%d/%m/%y', '%d-%m-%y', '%m/%d/%Y', '%m/%d/%y']:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    try:
        return pd.to_datetime(text, errors='coerce').to_pydatetime()
    except Exception:
        return None


def detect_header_row(raw_df: pd.DataFrame) -> int:
    for idx in range(min(raw_df.shape[0], 20)):
        row = raw_df.iloc[idx].fillna('')
        text_values = [normalize_header_name(v) for v in row.tolist()]
        score = 0
        for token in ['numruc', 'serie', 'documento', 'importe_cancelado', 'fecha_cancelacion', 'banco']:
            if token in text_values:
                score += 1
        if score >= 3:
            return idx
    raise ValueError('No se pudo localizar la fila de encabezado del reporte de ODE.')


def read_raw_ode_report(path: str | Path) -> pd.DataFrame:
    source_path = Path(path)
    if not source_path.exists():
        raise FileNotFoundError(f'No se encontró el archivo: {source_path}')

    try:
        df = pd.read_excel(source_path, header=None)
    except Exception:
        try:
            df = pd.read_excel(source_path, header=None, engine='xlrd')
        except Exception as exc:
            raise ValueError(f'No se pudo leer el archivo de origen: {exc}') from exc

    header_row = detect_header_row(df)
    headers = [normalize_header_name(cell) for cell in df.iloc[header_row].tolist()]
    data = df.iloc[header_row + 1:].copy()
    data.columns = headers
    data = data.dropna(how='all').reset_index(drop=True)

    if not data.empty:
        last_text = ' '.join([normalize_text(v) for v in data.iloc[-1].fillna('').tolist()])
        if 'TOTAL' in last_text.upper():
            data = data.iloc[:-1].copy()

    cleaned = data.copy()
    for column in cleaned.columns:
        cleaned[column] = cleaned[column].fillna('')

    required_aliases = {
        'tipo_movi': 'tipo_movi',
        'fecha_ingreso': 'fecha_ingreso',
        'grupo_cli': 'grupo_cli',
        'nombre_cli': 'nombre_cli',
        'numruc': 'numruc',
        'fecha_emi': 'fecha_emi',
        'serie': 'serie',
        'documento': 'documento',
        'kardex': 'kardex',
        'fecha_cancelacion': 'fecha_cancelacion',
        'banco': 'banco',
        'nro_cuenta_banco': 'nro_cuenta_banco',
        'movimiento': 'movimiento',
        'importe_notarial': 'importe_notarial',
        'importe_registral': 'importe_registral',
        'importe_cancelado': 'importe_cancelado',
        'saldo': 'saldo',
        'fecha_ven': 'fecha_ven',
        'monedatipo': 'monedatipo',
        'tipopago': 'tipopago',
        'codigoemp': 'codigoemp',
        'comentarios': 'comentarios',
    }
    rename_map = {key: value for key, value in required_aliases.items() if key in cleaned.columns}
    cleaned = cleaned.rename(columns=rename_map)

    for col in ['numruc', 'documento', 'serie', 'banco']:
        if col in cleaned.columns:
            if col in {'numruc', 'documento'}:
                cleaned[col] = cleaned[col].apply(normalize_document_value)
            elif col == 'serie':
                cleaned[col] = cleaned[col].apply(normalize_serie)
            elif col == 'banco':
                cleaned[col] = cleaned[col].apply(normalize_bank_name)

    for col in ['importe_cancelado', 'importe_notarial', 'importe_registral', 'saldo']:
        if col in cleaned.columns:
            cleaned[col] = cleaned[col].apply(normalize_amount)

    for col in ['fecha_cancelacion', 'fecha_emi', 'fecha_ven', 'fecha_ingreso']:
        if col in cleaned.columns:
            cleaned[col] = cleaned[col].apply(parse_date_value)

    cleaned = cleaned[cleaned.apply(lambda row: not row.astype(str).str.contains(r'^\s*$', regex=True).all(), axis=1)]
    return cleaned.reset_index(drop=True)


def split_ode_rows(raw_df: pd.DataFrame) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    raw = raw_df.copy()
    if raw.empty:
        return raw, raw.copy(), raw.copy()
    if 'serie' not in raw.columns:
        return raw, raw.copy(), raw.copy()
    raw['serie_norm'] = raw['serie'].astype(str).str.strip().str.upper().str.replace(' ', '')
    raw['serie_zfill'] = raw['serie_norm'].str.zfill(4)
    excluded = raw[raw['serie_zfill'].eq('0001')].copy()
    processed = raw[~raw.index.isin(excluded.index)].copy()
    return raw, excluded, processed


def load_ode_report(path: str | Path) -> pd.DataFrame:
    raw = read_raw_ode_report(path)
    _, _, processed = split_ode_rows(raw)
    return processed.reset_index(drop=True)


def get_month_end(periodo: str) -> datetime:
    if not periodo or len(periodo) != 6:
        raise ValueError(f'Periodo inválido: {periodo!r}. Debe tener formato YYYYMM.')
    period = pd.Period(periodo, freq='M')
    return period.end_time.normalize().to_pydatetime()


def get_period_from_report(df: pd.DataFrame) -> str:
    for column in ['fecha_cancelacion', 'fecha_emi']:
        if column in df.columns and df[column].notna().any():
            dates = pd.to_datetime(df[column].dropna(), errors='coerce')
            if not dates.empty:
                ts = dates.iloc[0]
                return f'{ts.year:04d}{ts.month:02d}'
    return datetime.now().strftime('%Y%m')


def validate_model_structure(model_path: Optional[str | Path], expected_columns: Optional[List[str]] = None) -> None:
    if not model_path:
        return
    expected = expected_columns or OUTPUT_COLUMNS
    path = Path(model_path)
    if not path.exists():
        raise FileNotFoundError(f'No se encontró el modelo: {path}')
    model = pd.read_excel(path, header=0)
    actual = list(model.columns)
    if actual != expected:
        raise ValueError(
            'La estructura del archivo modelo no coincide con la requerida. '
            f'Esperado: {expected}. Actual: {actual}'
        )


def auto_create_bank_config(config: Dict[str, Any], nombre_banco: str) -> Dict[str, Any]:
    config = config or {'bancos': {}}
    config.setdefault('bancos', {})
    bank_name = normalize_bank_name(nombre_banco)
    if not bank_name:
        return {'cuenta_contable': '10100011', 'medio_pago': '008'}
    if bank_name not in config['bancos']:
        config['bancos'][bank_name] = {'cuenta_contable': '10100011', 'medio_pago': '008'}
    return config['bancos'][bank_name]


def load_config(path: str | Path | None = None) -> Dict[str, Any]:
    config_path = Path(path) if path else get_app_dir() / 'config.json'
    if not config_path.exists():
        default = json.loads(json.dumps(DEFAULT_CONFIG))
        config_path.parent.mkdir(parents=True, exist_ok=True)
        with open(config_path, 'w', encoding='utf-8') as fh:
            json.dump(default, fh, ensure_ascii=False, indent=2)
        return default
    with open(config_path, 'r', encoding='utf-8') as fh:
        loaded = json.load(fh)
    merged = json.loads(json.dumps(DEFAULT_CONFIG))
    merged.update(loaded)
    merged['bancos'] = {**DEFAULT_CONFIG['bancos'], **loaded.get('bancos', {})}
    return merged


def save_config(path: str | Path, config: Dict[str, Any]) -> None:
    config_path = Path(path)
    config_path.parent.mkdir(parents=True, exist_ok=True)
    with open(config_path, 'w', encoding='utf-8') as fh:
        json.dump(config, fh, ensure_ascii=False, indent=2)


def validate_report(df: pd.DataFrame, periodo: Optional[str] = None, fecha_calculo: str = 'fecha_cancelacion') -> Tuple[List[str], List[str], List[str]]:
    errors: List[str] = []
    warnings: List[str] = []
    info: List[str] = []

    if df.empty:
        errors.append('El archivo no contiene movimientos válidos.')
        return errors, warnings, info

    required = ['numruc', 'serie', 'documento', 'importe_cancelado', 'banco']
    for col in required:
        if col not in df.columns:
            errors.append(f'Falta la columna requerida: {col}')

    if 'serie' in df.columns:
        empty_series = df[df['serie'].astype(str).str.strip() == '']
        if not empty_series.empty:
            errors.append('Hay filas con serie vacía.')

    if 'documento' in df.columns:
        empty_documents = df[df['documento'].astype(str).str.strip() == '']
        if not empty_documents.empty:
            errors.append('Hay filas con documento vacío.')

    if 'numruc' in df.columns:
        bad_ruc = df[df['numruc'].astype(str).str.len().lt(8)]
        if not bad_ruc.empty:
            errors.append('Hay filas con RUC demasiado corto o vacío.')

    if 'importe_cancelado' in df.columns:
        bad_amount = df[df['importe_cancelado'] <= 0]
        if not bad_amount.empty:
            errors.append('Hay movimientos con importe_cancelado vacío o menor o igual a cero.')

    if 'banco' in df.columns:
        empty_banks = df[df['banco'].astype(str).str.strip() == '']
        if not empty_banks.empty:
            errors.append('Hay movimientos sin nombre de banco. No se permite inferir un banco vacío.')

    if fecha_calculo == 'fecha_cancelacion':
        if 'fecha_cancelacion' not in df.columns:
            errors.append('Falta la columna fecha_cancelacion para el criterio seleccionado.')
        else:
            bad_dates = df[df['fecha_cancelacion'].isna()]
            if not bad_dates.empty:
                idxs = ', '.join(str(int(i + 1)) for i in bad_dates.index[:10])
                errors.append(f'Hay movimientos con fecha de cancelación inválida en filas: {idxs}.')
    elif fecha_calculo == 'fecha_emi':
        if 'fecha_emi' not in df.columns:
            errors.append('Falta la columna fecha_emi para el criterio seleccionado.')
        else:
            bad_dates = df[df['fecha_emi'].isna()]
            if not bad_dates.empty:
                idxs = ', '.join(str(int(i + 1)) for i in bad_dates.index[:10])
                errors.append(f'Hay movimientos con fecha de emisión inválida en filas: {idxs}.')
    elif fecha_calculo == 'ultimo_dia_mes':
        if not periodo:
            errors.append('Para ultimo_dia_mes es obligatorio indicar el periodo.')

    if 'serie' in df.columns:
        series_0001 = df[df['serie'].astype(str).str.zfill(4).eq('0001')]
        if not series_0001.empty:
            info.append(f'Se encontraron {len(series_0001)} registros de serie 0001 y fueron excluidos en la fase previa.')

    if 'banco' in df.columns and 'OTROS' in df['banco'].astype(str).unique():
        warnings.append('Los movimientos OTROS pueden corresponder a una retención del 3%. Verifique la cuenta contable configurada.')

    if not errors:
        info.append('Archivo legible y con estructura válida.')
    return errors, warnings, info


def build_carga_dataframe(processed_df: pd.DataFrame, periodo: str, config: Dict[str, Any], fecha_calculo: str = 'fecha_cancelacion', model_path: Optional[str | Path] = None) -> pd.DataFrame:
    config = {**DEFAULT_CONFIG, **(config or {})}
    config['bancos'] = {**DEFAULT_CONFIG['bancos'], **config.get('bancos', {})}
    if model_path:
        validate_model_structure(model_path, OUTPUT_COLUMNS)

    if processed_df.empty:
        return pd.DataFrame(columns=OUTPUT_COLUMNS)

    month_end = get_month_end(periodo)
    bank_order: List[str] = []
    seen: set[str] = set()
    for bank in BANK_PRIORITY:
        if bank in processed_df['banco'].astype(str).unique():
            bank_order.append(bank)
            seen.add(bank)
    bank_order.extend(sorted({b for b in processed_df['banco'].astype(str).unique() if b and b not in seen}, key=str.upper))

    rows: List[Dict[str, Any]] = []
    comprobante_index = 1
    for bank in bank_order:
        bank_rows = processed_df[processed_df['banco'].astype(str) == bank].copy()
        if bank_rows.empty:
            continue
        bank_cfg = config['bancos'].get(bank, {'cuenta_contable': '10100011', 'medio_pago': '008'})
        comprobante = f'{comprobante_index:04d}'

        rows.append({
            'CTA_CONTABLE': str(bank_cfg.get('cuenta_contable', '10100011')),
            'ANO_MES': str(periodo),
            'SUB_DIARIO': str(config.get('sub_diario', '01')),
            'COMPROBANTE': comprobante,
            'FEC_DOC': month_end,
            'TIP_ANEXO': str(config.get('tipo_anexo', '02')),
            'COD_ANEXO': '',
            'TIP_DOC': str(config.get('tipo_documento', '01')),
            'SERIE_NUM': '',
            'FEC_VENC': month_end,
            'MONEDA': str(config.get('moneda', 'MN')),
            'IMPORT_TOTAL': round(float(bank_rows['importe_cancelado'].sum()), 2),
            'TIP_CONEVR.': str(config.get('tipo_conversion', 'VTA')),
            'FEC_REG': month_end,
            'TIP_CAMBIO': '',
            'GLOSA': str(config.get('glosa', 'COBRANZA DEL MES')),
            'CENTRO_COSTOS': '',
            'GLOSA_MOV': str(config.get('glosa', 'COBRANZA DEL MES')),
            'DOC_ANULADO': str(config.get('doc_anulado', '0')),
            'DEBE HABER': 'D',
            'MEDIO_PAGO': str(bank_cfg.get('medio_pago', '008')),
            'NRO_FILE': '',
        })

        for _, row in bank_rows.iterrows():
            if fecha_calculo == 'fecha_cancelacion':
                date_value = row.get('fecha_cancelacion')
                if pd.isna(date_value):
                    raise ValueError(f'Falta fecha de cancelación para documento {row.get("documento", "N/A")} del banco {bank}.')
                doc_date = pd.to_datetime(date_value, errors='coerce').to_pydatetime()
            elif fecha_calculo == 'fecha_emi':
                date_value = row.get('fecha_emi')
                if pd.isna(date_value):
                    raise ValueError(f'Falta fecha de emisión para documento {row.get("documento", "N/A")} del banco {bank}.')
                doc_date = pd.to_datetime(date_value, errors='coerce').to_pydatetime()
            elif fecha_calculo == 'ultimo_dia_mes':
                doc_date = month_end
            else:
                doc_date = month_end

            rows.append({
                'CTA_CONTABLE': str(config.get('cuenta_clientes', '12120001')),
                'ANO_MES': str(periodo),
                'SUB_DIARIO': str(config.get('sub_diario', '01')),
                'COMPROBANTE': comprobante,
                'FEC_DOC': doc_date,
                'TIP_ANEXO': str(config.get('tipo_anexo', '02')),
                'COD_ANEXO': normalize_document_value(row.get('numruc', '')),
                'TIP_DOC': str(config.get('tipo_documento', '01')),
                'SERIE_NUM': serialize_serie_num({'serie': row.get('serie', ''), 'documento': row.get('documento', '')}),
                'FEC_VENC': doc_date,
                'MONEDA': str(config.get('moneda', 'MN')),
                'IMPORT_TOTAL': round(float(normalize_amount(row.get('importe_cancelado', 0))), 2),
                'TIP_CONEVR.': str(config.get('tipo_conversion', 'VTA')),
                'FEC_REG': month_end,
                'TIP_CAMBIO': '',
                'GLOSA': str(config.get('glosa', 'COBRANZA DEL MES')),
                'CENTRO_COSTOS': '',
                'GLOSA_MOV': str(config.get('glosa', 'COBRANZA DEL MES')),
                'DOC_ANULADO': str(config.get('doc_anulado', '0')),
                'DEBE HABER': 'H',
                'MEDIO_PAGO': str(bank_cfg.get('medio_pago', '008')),
                'NRO_FILE': '',
            })

        comprobante_index += 1

    output = pd.DataFrame(rows, columns=OUTPUT_COLUMNS)
    for column in ['COMPROBANTE', 'COD_ANEXO', 'SERIE_NUM', 'MONEDA', 'DOC_ANULADO', 'MEDIO_PAGO', 'ANO_MES']:
        output[column] = output[column].fillna('').astype(str)
    output = output.replace({np.nan: ''})
    output = output.fillna('')
    output['IMPORT_TOTAL'] = pd.to_numeric(output['IMPORT_TOTAL'], errors='coerce').fillna(0.0)
    return output


def build_conciliacion(carga: pd.DataFrame) -> List[Dict[str, Any]]:
    if carga.empty:
        return []
    rows: List[Dict[str, Any]] = []
    for comprobante, grupo in carga.groupby('COMPROBANTE', sort=False):
        banco = str(grupo['MEDIO_PAGO'].iloc[0]) if not grupo.empty else 'N/A'
        cantidad = int(len(grupo))
        debito = float(grupo[grupo['DEBE HABER'] == 'D']['IMPORT_TOTAL'].sum())
        credito = float(grupo[grupo['DEBE HABER'] == 'H']['IMPORT_TOTAL'].sum())
        diferencia = abs(debito - credito)
        estado = 'CUADRADO' if diferencia <= 0.01 else 'REVISAR'
        rows.append({
            'banco': banco,
            'comprobante': str(comprobante),
            'cantidad': cantidad,
            'debito': round(debito, 2),
            'credito': round(credito, 2),
            'diferencia': round(diferencia, 2),
            'estado': estado,
        })

    total_debito = float(carga[carga['DEBE HABER'] == 'D']['IMPORT_TOTAL'].sum())
    total_credito = float(carga[carga['DEBE HABER'] == 'H']['IMPORT_TOTAL'].sum())
    rows.append({
        'banco': 'TOTAL',
        'comprobante': 'TOTAL',
        'cantidad': int(len(carga)),
        'debito': round(total_debito, 2),
        'credito': round(total_credito, 2),
        'diferencia': round(abs(total_debito - total_credito), 2),
        'estado': 'CUADRADO' if abs(total_debito - total_credito) <= 0.01 else 'REVISAR',
    })
    return rows


def create_validation_report(summary: Dict[str, Any], source_path: Path, output_path: Path) -> None:
    from openpyxl.styles import Font, PatternFill

    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        summary_df = pd.DataFrame([
            ['Fecha de generación', summary['fecha_generacion']],
            ['Archivo de origen', str(source_path)],
            ['Periodo', summary['periodo']],
            ['Cantidad original', summary['cantidad_original']],
            ['Cantidad excluida', summary['cantidad_excluida']],
            ['Cantidad procesada', summary['cantidad_procesada']],
            ['Total procesado', summary['total_procesado']],
            ['Total débito', summary['total_debito']],
            ['Total crédito', summary['total_credito']],
            ['Diferencia', summary['diferencia']],
            ['Resultado', summary['resultado']],
        ], columns=['Campo', 'Valor'])
        summary_df.to_excel(writer, index=False, sheet_name='Resumen')

        processed_df = pd.DataFrame(summary.get('procesados', []))
        if not processed_df.empty:
            processed_df.to_excel(writer, index=False, sheet_name='Procesados')
        else:
            pd.DataFrame([{'mensaje': 'Sin registros procesados'}]).to_excel(writer, index=False, sheet_name='Procesados')

        excluidos_df = pd.DataFrame(summary.get('excluidos', []))
        if not excluidos_df.empty:
            excluidos_df.to_excel(writer, index=False, sheet_name='Excluidos')
        else:
            pd.DataFrame([{'mensaje': 'Sin exclusiones'}]).to_excel(writer, index=False, sheet_name='Excluidos')

        observ_df = pd.DataFrame(summary.get('observaciones', []))
        if not observ_df.empty:
            observ_df.to_excel(writer, index=False, sheet_name='Observaciones')
        else:
            pd.DataFrame([{'mensaje': 'Sin observaciones'}]).to_excel(writer, index=False, sheet_name='Observaciones')

        conciliacion_df = pd.DataFrame(summary.get('conciliacion', []))
        if not conciliacion_df.empty:
            conciliacion_df.to_excel(writer, index=False, sheet_name='Conciliacion')
        else:
            pd.DataFrame([{'banco': 'N/A', 'comprobante': 'N/A', 'cantidad': 0, 'debito': 0, 'credito': 0, 'diferencia': 0, 'estado': 'N/A'}]).to_excel(writer, index=False, sheet_name='Conciliacion')

        fill = PatternFill('solid', fgColor='1F2A44')
        font = Font(color='FFFFFF', bold=True)
        for ws in writer.sheets.values():
            ws.freeze_panes = 'A2'
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = fill
                cell.font = font


def export_carga_workbook(output_df: pd.DataFrame, output_path: Path, file_label: str = 'CARGA') -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = None
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = 'Hoja1'
        ws.freeze_panes = 'A2'
        ws.append(list(OUTPUT_COLUMNS))
        for _, row in output_df.iterrows():
            values = []
            for col in OUTPUT_COLUMNS:
                value = row.get(col, '')
                if value is None or (isinstance(value, float) and pd.isna(value)):
                    values.append('')
                else:
                    values.append(value)
            ws.append(values)
        header_fill = PatternFill(fill_type='solid', fgColor='1F2A44')
        header_font = Font(color='FFFFFF', bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        for col in ws.columns:
            max_len = 0
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col[0].column_letter].width = min(max(12, max_len + 2), 22)
        ws.auto_filter.ref = ws.dimensions
        wb.save(output_path)
    except Exception:
        if wb is not None:
            try:
                wb.save(output_path)
            except Exception:
                pass
        raise


def ensure_no_nan(output_df: pd.DataFrame) -> pd.DataFrame:
    cleaned = output_df.copy()
    for column in cleaned.columns:
        cleaned[column] = cleaned[column].apply(lambda value: '' if value is None or (isinstance(value, float) and pd.isna(value)) or str(value).lower() in {'nan', 'nat'} else value)
    return cleaned


def process_report(source_path: str | Path, output_dir: str | Path, config: Optional[Dict[str, Any]] = None, fecha_calculo: str = 'fecha_cancelacion', periodo: Optional[str] = None, model_path: Optional[str | Path] = None) -> Dict[str, Any]:
    if config is None:
        config = load_config(get_app_dir() / 'config.json')
    source_path = Path(source_path)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    raw = read_raw_ode_report(source_path)
    if raw.empty:
        raise ValueError('No se encontraron movimientos válidos en el archivo de ODE.')

    raw, excluded, processed = split_ode_rows(raw)
    if periodo is None:
        periodo = get_period_from_report(processed) or datetime.now().strftime('%Y%m')

    if model_path is not None:
        validate_model_structure(model_path, OUTPUT_COLUMNS)

    errors, warnings, info = validate_report(processed, periodo=periodo, fecha_calculo=fecha_calculo)
    if errors:
        raise ValueError('Se encontraron errores bloqueantes:\n- ' + '\n- '.join(errors))

    for idx in processed.index:
        if pd.isna(processed.at[idx, 'banco']) or str(processed.at[idx, 'banco']).strip() == '':
            raise ValueError(f'Banco vacío en la fila {idx + 1}. No se puede inferir el medio de cobranza.')
        normalized = normalize_bank_name(processed.at[idx, 'banco'])
        auto_create_bank_config(config, normalized)

    carga = build_carga_dataframe(processed, periodo, config, fecha_calculo=fecha_calculo, model_path=model_path)
    carga = ensure_no_nan(carga)

    if carga.empty:
        raise ValueError('No se pudieron generar registros contables a partir del reporte.')

    total_procesado = float(processed['importe_cancelado'].sum()) if 'importe_cancelado' in processed.columns else 0.0
    total_debito = float(carga[carga['DEBE HABER'] == 'D']['IMPORT_TOTAL'].sum()) if 'DEBE HABER' in carga.columns else 0.0
    total_credito = float(carga[carga['DEBE HABER'] == 'H']['IMPORT_TOTAL'].sum()) if 'DEBE HABER' in carga.columns else 0.0
    diff = abs(total_debito - total_credito)

    if not np.isclose(total_debito, total_credito, atol=0.01):
        raise ValueError('La carga generada no cuadra: debe y crédito deben coincidir.')
    if not np.isclose(total_procesado, total_debito, atol=0.01):
        raise ValueError('El total procesado no coincide con el débito generado.')
    if not np.isclose(total_procesado, total_credito, atol=0.01):
        raise ValueError('El total procesado no coincide con el crédito generado.')
    movimiento = carga[carga['DEBE HABER'] == 'H']
    if not movimiento.empty:
        if movimiento['SERIE_NUM'].astype(str).str.strip().eq('').any():
            raise ValueError('Hay filas de movimiento con SERIE_NUM vacío.')
        if movimiento['COD_ANEXO'].astype(str).str.strip().eq('').any():
            raise ValueError('Hay filas de movimiento con COD_ANEXO vacío.')
        if (movimiento['IMPORT_TOTAL'] <= 0).any():
            raise ValueError('Hay importes menores o iguales a cero en movimientos de crédito.')
    if (carga[carga['DEBE HABER'] == 'D']['IMPORT_TOTAL'] <= 0).any():
        raise ValueError('Hay importes menores o iguales a cero en los totales de banco.')
    if carga.isna().values.any():
        raise ValueError('Existen celdas NaN/NaT/None en la carga.')

    conciliacion = build_conciliacion(carga)
    summary = {
        'fecha_generacion': datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
        'periodo': periodo,
        'cantidad_original': len(raw),
        'cantidad_excluida': len(excluded),
        'cantidad_procesada': len(processed),
        'total_procesado': round(total_procesado, 2),
        'total_debito': round(total_debito, 2),
        'total_credito': round(total_credito, 2),
        'diferencia': round(diff, 2),
        'resultado': 'CUADRADO' if diff <= 0.01 else 'REVISAR',
        'procesados': processed.to_dict(orient='records'),
        'excluidos': excluded.to_dict(orient='records'),
        'conciliacion': conciliacion,
        'observaciones': [{'tipo': 'warning', 'mensaje': m} for m in warnings] + [{'tipo': 'info', 'mensaje': m} for m in info],
    }

    out_file = output_dir / f'CARGA_COBRANZAS_ODE_{periodo}.xlsx'
    report_file = output_dir / f'REPORTE_VALIDACION_ODE_{periodo}.xlsx'
    if out_file.exists():
        out_file = output_dir / f'CARGA_COBRANZAS_ODE_{periodo}_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'
    if report_file.exists():
        report_file = output_dir / f'REPORTE_VALIDACION_ODE_{periodo}_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'

    export_carga_workbook(carga, out_file, 'CARGA')
    create_validation_report(summary, source_path, report_file)
    result = {
        'summary': {
            'archivo_origen': str(source_path),
            'periodo': periodo,
            'cantidad_original': len(raw),
            'cantidad_excluida': len(excluded),
            'cantidad_procesada': len(processed),
            'total_procesado': round(total_procesado, 2),
            'total_debito': round(total_debito, 2),
            'total_credito': round(total_credito, 2),
            'diferencia': round(diff, 2),
            'resultado': 'CUADRADO' if diff <= 0.01 else 'REVISAR',
            'conciliacion': conciliacion,
        },
        'output_file': str(out_file),
        'report_file': str(report_file),
        'processed_df': processed,
        'excluded_df': excluded,
        'output_df': carga,
        'errors': errors,
        'warnings': warnings,
        'info': info,
        'period': periodo,
    }
    return result
